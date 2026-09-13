"""
generate_zoro_tax_refund_excels.py

Generates executive-grade, beautifully styled Excel workbooks for submitting
sales tax refund claims to Zoro's Accounting Department.

Covered Accounts (Operating under M&R Market, LLC):
  1. inuhrecomllc   | Customer #25217105 | inuhrecomllc5@gmail.com
  2. salesmercentile| Customer #25017919 | salesmercentile@gmail.com
  3. salesmercentile1| Customer #25061417 | salesmercentil1@gmail.com

Output Files:
  - Automations/data/Zoro_Sales_Tax_Refund_Master_MR_Market_LLC.xlsx
  - Automations/data/Zoro_Sales_Tax_Refund_inuhrecomllc_CUST25217105.xlsx
  - Automations/data/Zoro_Sales_Tax_Refund_salesmercentile_CUST25017919.xlsx
  - Automations/data/Zoro_Sales_Tax_Refund_salesmercentile1_CUST25061417.xlsx
"""

import os
import sqlite3
import pandas as pd
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------
# CONSTANTS & CONFIGURATION
# ---------------------------------------------------------
CUTOFF_6M = "2026-03-11"  # 6 months prior to Sep 11, 2026
CURRENT_DATE = "2026-09-11"

# 44 states officially confirmed exempt by Zoro Tax Exemption Specialist Kelly (Sep 11, 2026)
KELLY_APPROVED_STATES = {
    "AL", "AK", "AZ", "AR", "CT", "CA", "CO", "GA", "FL", "IN", "IA", "ID",
    "IL", "KS", "KY", "LA", "MI", "ME", "MO", "MN", "MD", "MA", "NE", "NV",
    "NJ", "NY", "NC", "ND", "NM", "OH", "OK", "PA", "RI", "SD", "SC", "TN",
    "TX", "UT", "VT", "VA", "WA", "WI", "WV", "WY"
}

STATE_NAMES = {
    "AL": "Alabama", "AK": "Alaska", "AZ": "Arizona", "AR": "Arkansas", "CA": "California",
    "CO": "Colorado", "CT": "Connecticut", "DE": "Delaware", "DC": "District of Columbia",
    "FL": "Florida", "GA": "Georgia", "ID": "Idaho", "IL": "Illinois", "IN": "Indiana",
    "IA": "Iowa", "KS": "Kansas", "KY": "Kentucky", "LA": "Louisiana", "ME": "Maine",
    "MD": "Maryland", "MA": "Massachusetts", "MI": "Michigan", "MN": "Minnesota",
    "MS": "Mississippi", "MO": "Missouri", "MT": "Montana", "NE": "Nebraska", "NV": "Nevada",
    "NH": "New Hampshire", "NJ": "New Jersey", "NM": "New Mexico", "NY": "New York",
    "NC": "North Carolina", "ND": "North Dakota", "OH": "Ohio", "OK": "Oklahoma",
    "OR": "Oregon", "PA": "Pennsylvania", "RI": "Rhode Island", "SC": "South Carolina",
    "SD": "South Dakota", "TN": "Tennessee", "TX": "Texas", "UT": "Utah", "VT": "Vermont",
    "VA": "Virginia", "WA": "Washington", "WV": "West Virginia", "WI": "Wisconsin", "WY": "Wyoming"
}

ACCOUNTS_META = {
    "inhur_llc": {
        "cust_id": "25217105",
        "cust_str": "CUST25217105",
        "name": "inuhrecomllc",
        "email": "inuhrecomllc5@gmail.com",
        "entity": "M&R Market, LLC",
        "tab_name": "inuhrecomllc (CUST25217105)",
    },
    "salesmercentile": {
        "cust_id": "25017919",
        "cust_str": "CUST25017919",
        "name": "salesmercentile",
        "email": "salesmercentile@gmail.com",
        "entity": "M&R Market, LLC",
        "tab_name": "salesmercentile (CUST25017919)",
    },
    "salesmercentile1": {
        "cust_id": "25061417",
        "cust_str": "CUST25061417",
        "name": "salesmercentile1",
        "email": "salesmercentil1@gmail.com",
        "entity": "M&R Market, LLC",
        "tab_name": "salesmercentile1 (CUST25061417)",
    },
}

# ---------------------------------------------------------
# STYLING PALETTE
# ---------------------------------------------------------
FONT_FAMILY = "Segoe UI"

FONT_TITLE = Font(name=FONT_FAMILY, size=15, bold=True, color="FFFFFF")
FONT_SUBTITLE = Font(name=FONT_FAMILY, size=10, italic=True, color="E2E8F0")
FONT_SECTION = Font(name=FONT_FAMILY, size=12, bold=True, color="1E3A8A")
FONT_TH = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
FONT_TD = Font(name=FONT_FAMILY, size=9.5, color="1E293B")
FONT_TD_BOLD = Font(name=FONT_FAMILY, size=9.5, bold=True, color="1E293B")
FONT_TOTAL = Font(name=FONT_FAMILY, size=10, bold=True, color="0F172A")
FONT_TAX_HIGHLIGHT = Font(name=FONT_FAMILY, size=9.5, bold=True, color="15803D")

# Fills
FILL_TITLE = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")       # Dark Royal Navy
FILL_TH = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")          # Deep Navy
FILL_TH_SUB = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")      # Royal Blue
FILL_ZEBRA = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")       # Soft Off-white
FILL_WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
FILL_TOTAL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")       # Soft Slate
FILL_CARD = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
FILL_TAX_CELL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")    # Soft Mint Green
FILL_ALERT_CELL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Soft Amber

# Borders
BORDER_THIN_GRAY = Side(style="thin", color="CBD5E1")
BORDER_MEDIUM_NAVY = Side(style="medium", color="1E3A8A")
BORDER_DOUBLE_NAVY = Side(style="double", color="1E3A8A")

BORDER_CELL = Border(left=BORDER_THIN_GRAY, right=BORDER_THIN_GRAY, top=BORDER_THIN_GRAY, bottom=BORDER_THIN_GRAY)
BORDER_HEADER = Border(left=BORDER_THIN_GRAY, right=BORDER_THIN_GRAY, top=BORDER_MEDIUM_NAVY, bottom=BORDER_MEDIUM_NAVY)
BORDER_TOTAL_ROW = Border(left=BORDER_THIN_GRAY, right=BORDER_THIN_GRAY, top=BORDER_THIN_GRAY, bottom=BORDER_DOUBLE_NAVY)

# Alignments
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Formats
FMT_CURRENCY = "$#,##0.00"
FMT_INT = "#,##0"
FMT_PCT = "0.0%"

# ---------------------------------------------------------
# DATA FETCHING & PREPARATION
# ---------------------------------------------------------
def load_tax_data(db_path):
    conn = sqlite3.connect(db_path)
    q = """
        SELECT 
            wb_order_id,
            gmail_account,
            order_date,
            email_timestamp,
            subtotal,
            tax,
            order_total,
            ship_city,
            ship_state,
            ship_zip
        FROM archive_zoro
        WHERE gmail_account IN ('inhur_llc', 'salesmercentile', 'salesmercentile1')
          AND tax > 0
        ORDER BY email_timestamp DESC
    """
    df = pd.read_sql_query(q, conn)
    conn.close()

    # Parse and normalize dates
    df["dt_order"] = pd.to_datetime(df["order_date"], format="mixed", errors="coerce")
    df["dt_email"] = pd.to_datetime(df["email_timestamp"], errors="coerce")
    
    # Clean order date string (YYYY-MM-DD)
    df["order_date_clean"] = df["dt_order"].dt.strftime("%Y-%m-%d")
    # Fallback to email date if missing
    df["order_date_clean"] = df["order_date_clean"].fillna(df["dt_email"].dt.strftime("%Y-%m-%d"))
    
    # State cleaning
    df["ship_state_clean"] = df["ship_state"].fillna("").astype(str).str.strip().str.upper()
    df["state_name"] = df["ship_state_clean"].map(STATE_NAMES).fillna(df["ship_state_clean"])
    
    # City & Zip
    df["ship_city_clean"] = df["ship_city"].fillna("").astype(str).str.strip().str.title()
    df["ship_zip_clean"] = df["ship_zip"].fillna("").astype(str).str.strip()
    
    # Numbers
    df["subtotal"] = pd.to_numeric(df["subtotal"], errors="coerce").fillna(0.0)
    df["tax"] = pd.to_numeric(df["tax"], errors="coerce").fillna(0.0)
    df["order_total"] = pd.to_numeric(df["order_total"], errors="coerce").fillna(0.0)

    # 6-Month Accounting Window flag (>= 2026-03-11)
    df["is_6m"] = (df["dt_email"] >= CUTOFF_6M) | (df["dt_order"] >= CUTOFF_6M)
    
    # Exemption flag
    df["is_approved_state"] = df["ship_state_clean"].isin(KELLY_APPROVED_STATES)
    df["exemption_status"] = df["is_approved_state"].apply(
        lambda x: "Approved Resale State" if x else "State Pending Exemption"
    )
    df["refund_eligibility"] = df["is_6m"].apply(
        lambda x: "Eligible (6-Month Zoro Accounting Window)" if x else "Historical (> 6 Months - State Refund Claim)"
    )

    # Account metadata enrichment
    df["cust_id"] = df["gmail_account"].apply(lambda a: ACCOUNTS_META[a]["cust_id"])
    df["cust_str"] = df["gmail_account"].apply(lambda a: ACCOUNTS_META[a]["cust_str"])
    df["account_name"] = df["gmail_account"].apply(lambda a: ACCOUNTS_META[a]["name"])
    df["account_email"] = df["gmail_account"].apply(lambda a: ACCOUNTS_META[a]["email"])
    df["entity_name"] = df["gmail_account"].apply(lambda a: ACCOUNTS_META[a]["entity"])

    return df

# ---------------------------------------------------------
# EXCEL GENERATION HELPERS
# ---------------------------------------------------------
def create_title_banner(ws, title_text, subtitle_text, max_col=12):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    cell1 = ws.cell(row=1, column=1, value=title_text)
    cell1.font = FONT_TITLE
    cell1.fill = FILL_TITLE
    cell1.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    cell2 = ws.cell(row=2, column=1, value=subtitle_text)
    cell2.font = FONT_SUBTITLE
    cell2.fill = FILL_TITLE
    cell2.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    ws.row_dimensions[3].height = 10  # blank spacer row


def style_table_header(ws, row_idx, headers, fill=FILL_TH):
    ws.row_dimensions[row_idx].height = 24
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=row_idx, column=col_idx, value=h)
        c.font = FONT_TH
        c.fill = fill
        c.alignment = ALIGN_HEADER
        c.border = BORDER_HEADER


def auto_fit_columns(ws, max_len_cap=45, min_len=11):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            # If cell is merged and not top-left, skip
            if cell.coordinate in ws.merged_cells:
                continue
            if len(val_str) > max_len and "\n" not in val_str:
                max_len = len(val_str)
        col_width = min(max_len_cap, max(min_len, max_len + 3))
        ws.column_dimensions[col_letter].width = col_width


def add_kpi_card(ws, start_row, start_col, end_col, title, value_str, note, val_color="1E3A8A", bg_color="F1F5F9"):
    # Merge cells for card
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=end_col)
    c_title = ws.cell(row=start_row, column=start_col, value=title.upper())
    c_title.font = Font(name=FONT_FAMILY, size=8.5, bold=True, color="64748B")
    c_title.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=start_row+1, start_column=start_col, end_row=start_row+1, end_column=end_col)
    c_val = ws.cell(row=start_row+1, column=start_col, value=value_str)
    c_val.font = Font(name=FONT_FAMILY, size=15, bold=True, color=val_color)
    c_val.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=start_row+2, start_column=start_col, end_row=start_row+2, end_column=end_col)
    c_note = ws.cell(row=start_row+2, column=start_col, value=note)
    c_note.font = Font(name=FONT_FAMILY, size=8, italic=True, color="64748B")
    c_note.alignment = Alignment(horizontal="center", vertical="center")

    card_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    for r in range(start_row, start_row + 3):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = card_fill
            cell.border = BORDER_CELL
    
    ws.row_dimensions[start_row].height = 15
    ws.row_dimensions[start_row+1].height = 24
    ws.row_dimensions[start_row+2].height = 14


def write_order_table(ws, df_orders, start_row=4, is_single_account=False):
    """Writes standard itemized order table and returns ending row."""
    if is_single_account:
        headers = [
            "Zoro Order #", "Order Date", "Ship State", "Ship City", "Ship Zip",
            "Item Subtotal ($)", "Sales Tax Charged ($)", "Order Total ($)",
            "Approved Resale State?", "Accounting Refund Eligibility"
        ]
    else:
        headers = [
            "Customer #", "Account Name", "Account Email", "Zoro Order #",
            "Order Date", "Ship State", "Ship City", "Ship Zip",
            "Item Subtotal ($)", "Sales Tax Charged ($)", "Order Total ($)",
            "Approved Resale State?", "Accounting Refund Eligibility"
        ]

    style_table_header(ws, start_row, headers)
    
    row_idx = start_row + 1
    for _, r in df_orders.iterrows():
        is_even = ((row_idx - start_row) % 2 == 0)
        row_fill = FILL_ZEBRA if is_even else FILL_WHITE
        ws.row_dimensions[row_idx].height = 18

        if is_single_account:
            vals = [
                (r["wb_order_id"], ALIGN_CENTER, None, FONT_TD_BOLD),
                (r["order_date_clean"], ALIGN_CENTER, None, FONT_TD),
                (r["ship_state_clean"], ALIGN_CENTER, None, FONT_TD_BOLD),
                (r["ship_city_clean"], ALIGN_LEFT, None, FONT_TD),
                (r["ship_zip_clean"], ALIGN_CENTER, None, FONT_TD),
                (r["subtotal"], ALIGN_RIGHT, FMT_CURRENCY, FONT_TD),
                (r["tax"], ALIGN_RIGHT, FMT_CURRENCY, FONT_TAX_HIGHLIGHT),
                (r["order_total"], ALIGN_RIGHT, FMT_CURRENCY, FONT_TD),
                (r["exemption_status"], ALIGN_CENTER, None, FONT_TD),
                (r["refund_eligibility"], ALIGN_CENTER, None, FONT_TD),
            ]
        else:
            vals = [
                (r["cust_id"], ALIGN_CENTER, None, FONT_TD),
                (r["account_name"], ALIGN_LEFT, None, FONT_TD),
                (r["account_email"], ALIGN_LEFT, None, FONT_TD),
                (r["wb_order_id"], ALIGN_CENTER, None, FONT_TD_BOLD),
                (r["order_date_clean"], ALIGN_CENTER, None, FONT_TD),
                (r["ship_state_clean"], ALIGN_CENTER, None, FONT_TD_BOLD),
                (r["ship_city_clean"], ALIGN_LEFT, None, FONT_TD),
                (r["ship_zip_clean"], ALIGN_CENTER, None, FONT_TD),
                (r["subtotal"], ALIGN_RIGHT, FMT_CURRENCY, FONT_TD),
                (r["tax"], ALIGN_RIGHT, FMT_CURRENCY, FONT_TAX_HIGHLIGHT),
                (r["order_total"], ALIGN_RIGHT, FMT_CURRENCY, FONT_TD),
                (r["exemption_status"], ALIGN_CENTER, None, FONT_TD),
                (r["refund_eligibility"], ALIGN_CENTER, None, FONT_TD),
            ]

        for col_idx, (val, align, num_fmt, font) in enumerate(vals, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.alignment = align
            c.font = font
            c.border = BORDER_CELL
            # If tax column, use mint fill
            tax_col = 7 if is_single_account else 10
            c.fill = FILL_TAX_CELL if col_idx == tax_col else row_fill
            if num_fmt:
                c.number_format = num_fmt

        row_idx += 1

    # Totals Row
    ws.row_dimensions[row_idx].height = 22
    total_cols = len(headers)
    
    # Label
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=(5 if is_single_account else 8))
    c_tot_label = ws.cell(row=row_idx, column=1, value="TOTAL CLAIM AMOUNT")
    c_tot_label.font = FONT_TOTAL
    c_tot_label.alignment = Alignment(horizontal="right", vertical="center")
    c_tot_label.fill = FILL_TOTAL
    for col in range(1, (6 if is_single_account else 9)):
        ws.cell(row=row_idx, column=col).fill = FILL_TOTAL
        ws.cell(row=row_idx, column=col).border = BORDER_TOTAL_ROW

    # Subtotal formula
    subtotal_col = 6 if is_single_account else 9
    tax_col = 7 if is_single_account else 10
    total_col = 8 if is_single_account else 11

    c_sub = ws.cell(row=row_idx, column=subtotal_col, value=f"=SUM({get_column_letter(subtotal_col)}{start_row+1}:{get_column_letter(subtotal_col)}{row_idx-1})")
    c_sub.font = FONT_TOTAL
    c_sub.alignment = ALIGN_RIGHT
    c_sub.number_format = FMT_CURRENCY
    c_sub.fill = FILL_TOTAL
    c_sub.border = BORDER_TOTAL_ROW

    c_tax = ws.cell(row=row_idx, column=tax_col, value=f"=SUM({get_column_letter(tax_col)}{start_row+1}:{get_column_letter(tax_col)}{row_idx-1})")
    c_tax.font = Font(name=FONT_FAMILY, size=10, bold=True, color="15803D")
    c_tax.alignment = ALIGN_RIGHT
    c_tax.number_format = FMT_CURRENCY
    c_tax.fill = FILL_TAX_CELL
    c_tax.border = BORDER_TOTAL_ROW

    c_tot = ws.cell(row=row_idx, column=total_col, value=f"=SUM({get_column_letter(total_col)}{start_row+1}:{get_column_letter(total_col)}{row_idx-1})")
    c_tot.font = FONT_TOTAL
    c_tot.alignment = ALIGN_RIGHT
    c_tot.number_format = FMT_CURRENCY
    c_tot.fill = FILL_TOTAL
    c_tot.border = BORDER_TOTAL_ROW

    # Remaining empty cells in total row
    for col in range(total_col + 1, total_cols + 1):
        c_rem = ws.cell(row=row_idx, column=col, value="")
        c_rem.fill = FILL_TOTAL
        c_rem.border = BORDER_TOTAL_ROW

    # Enable auto filter and freeze panes
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(total_cols)}{row_idx}"
    ws.freeze_panes = f"A{start_row+1}"
    ws.views.sheetView[0].showGridLines = True
    auto_fit_columns(ws)

    return row_idx


# ---------------------------------------------------------
# MASTER WORKBOOK BUILDER
# ---------------------------------------------------------
def generate_master_workbook(df, out_path):
    print(f"Generating Master Workbook: {out_path}...")
    wb = openpyxl.Workbook()
    # Default sheet
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"

    df_6m = df[df["is_6m"]].copy()
    df_hist = df[~df["is_6m"]].copy()

    total_6m_tax = df_6m["tax"].sum()
    total_6m_orders = len(df_6m)
    total_6m_spend = df_6m["order_total"].sum()
    total_all_tax = df["tax"].sum()

    # =====================================================
    # TAB 1: EXECUTIVE SUMMARY
    # =====================================================
    create_title_banner(
        ws_summary,
        "M&R MARKET, LLC — SALES TAX EXEMPTION REFUND CLAIM",
        "Formal Multi-Account Request for Zoro Accounting Department | 44 States Resale Exemption Review",
        max_col=10
    )

    # Formal Letterhead Metadata
    meta_rows = [
        ("Business Entity:", "M&R Market, LLC (Owner: Mark Booth)"),
        ("Primary Contact:", "salesmercentile@gmail.com | Phone: 855.BUY.ZORO"),
        ("Zoro Specialist:", "Kelly — Zoro Tax Exemption Specialist (salestaxzoro@zoro.com)"),
        ("Confirmation Date:", f"{CURRENT_DATE} (Approval of 44 Resale States across all 3 accounts)"),
        ("Accounting Period:", f"6-Month Direct Refund Window: March 11, 2026 – September 11, 2026"),
    ]
    for idx, (label, val) in enumerate(meta_rows, start=4):
        ws_summary.row_dimensions[idx].height = 18
        c_lbl = ws_summary.cell(row=idx, column=1, value=label)
        c_lbl.font = FONT_TD_BOLD
        c_lbl.alignment = ALIGN_LEFT
        ws_summary.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=7)
        c_val = ws_summary.cell(row=idx, column=2, value=val)
        c_val.font = FONT_TD
        c_val.alignment = ALIGN_LEFT

    ws_summary.row_dimensions[9].height = 10

    # 4 KPI Stat Cards (Rows 10 - 12)
    add_kpi_card(ws_summary, start_row=10, start_col=1, end_col=2,
                 title="6-Month Tax Refund Requested",
                 value_str=f"${total_6m_tax:,.2f}",
                 note="Direct Zoro Credit/Check",
                 val_color="15803D", bg_color="DCFCE7")

    add_kpi_card(ws_summary, start_row=10, start_col=3, end_col=4,
                 title="Eligible 6-Month Orders",
                 value_str=f"{total_6m_orders:,}",
                 note="All Starting with 'WB'",
                 val_color="1D4ED8", bg_color="EFF6FF")

    add_kpi_card(ws_summary, start_row=10, start_col=5, end_col=7,
                 title="Total Zoro Purchasing Volume",
                 value_str=f"${total_6m_spend:,.2f}",
                 note="March 11 – September 11, 2026",
                 val_color="1E293B", bg_color="F1F5F9")

    add_kpi_card(ws_summary, start_row=10, start_col=8, end_col=10,
                 title="Approved Resale States",
                 value_str="44 States Approved",
                 note="Kelly Confirmed (Sep 11, 2026)",
                 val_color="B45309", bg_color="FEF3C7")

    ws_summary.row_dimensions[13].height = 14

    # Table 1: Account Summary Table
    ws_summary.cell(row=14, column=1, value="1. RECONCILIATION BY CUSTOMER ACCOUNT (6-MONTH ACCOUNTING ELIGIBILITY)").font = FONT_SECTION
    ws_summary.row_dimensions[14].height = 20

    acct_headers = [
        "Business Entity", "Account Name", "Zoro Customer #", "Account Email",
        "6-Mo Orders", "6-Mo Tax Refund ($)", "6-Mo Spend ($)",
        "Pre-6Mo Orders", "Pre-6Mo Tax ($)", "Total All-Time Tax ($)"
    ]
    style_table_header(ws_summary, 15, acct_headers)

    curr_row = 16
    for acct_key in ["inhur_llc", "salesmercentile", "salesmercentile1"]:
        meta = ACCOUNTS_META[acct_key]
        df_a_6m = df_6m[df_6m["gmail_account"] == acct_key]
        df_a_hist = df_hist[df_hist["gmail_account"] == acct_key]
        df_a_all = df[df["gmail_account"] == acct_key]

        ws_summary.row_dimensions[curr_row].height = 20
        row_fill = FILL_ZEBRA if curr_row % 2 == 0 else FILL_WHITE

        vals = [
            (meta["entity"], ALIGN_LEFT, None, FONT_TD),
            (meta["name"], ALIGN_LEFT, None, FONT_TD_BOLD),
            (meta["cust_str"], ALIGN_CENTER, None, FONT_TD_BOLD),
            (meta["email"], ALIGN_LEFT, None, FONT_TD),
            (len(df_a_6m), ALIGN_RIGHT, FMT_INT, FONT_TD),
            (df_a_6m["tax"].sum(), ALIGN_RIGHT, FMT_CURRENCY, FONT_TAX_HIGHLIGHT),
            (df_a_6m["order_total"].sum(), ALIGN_RIGHT, FMT_CURRENCY, FONT_TD),
            (len(df_a_hist), ALIGN_RIGHT, FMT_INT, FONT_TD),
            (df_a_hist["tax"].sum(), ALIGN_RIGHT, FMT_CURRENCY, FONT_TD),
            (df_a_all["tax"].sum(), ALIGN_RIGHT, FMT_CURRENCY, FONT_TD_BOLD),
        ]

        for col_idx, (val, align, num_fmt, font) in enumerate(vals, start=1):
            c = ws_summary.cell(row=curr_row, column=col_idx, value=val)
            c.alignment = align
            c.font = font
            c.border = BORDER_CELL
            c.fill = FILL_TAX_CELL if col_idx == 6 else row_fill
            if num_fmt:
                c.number_format = num_fmt

        curr_row += 1

    # Total Row for Account Table
    ws_summary.row_dimensions[curr_row].height = 22
    ws_summary.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=4)
    c_tot = ws_summary.cell(row=curr_row, column=1, value="TOTAL M&R MARKET, LLC")
    c_tot.font = FONT_TOTAL
    c_tot.alignment = Alignment(horizontal="right", vertical="center")
    for col in range(1, 5):
        ws_summary.cell(row=curr_row, column=col).fill = FILL_TOTAL
        ws_summary.cell(row=curr_row, column=col).border = BORDER_TOTAL_ROW

    tot_cols = [
        (5, f"=SUM(E16:E{curr_row-1})", FMT_INT, FONT_TOTAL, FILL_TOTAL),
        (6, f"=SUM(F16:F{curr_row-1})", FMT_CURRENCY, Font(name=FONT_FAMILY, size=10, bold=True, color="15803D"), FILL_TAX_CELL),
        (7, f"=SUM(G16:G{curr_row-1})", FMT_CURRENCY, FONT_TOTAL, FILL_TOTAL),
        (8, f"=SUM(H16:H{curr_row-1})", FMT_INT, FONT_TOTAL, FILL_TOTAL),
        (9, f"=SUM(I16:I{curr_row-1})", FMT_CURRENCY, FONT_TOTAL, FILL_TOTAL),
        (10, f"=SUM(J16:J{curr_row-1})", FMT_CURRENCY, FONT_TOTAL, FILL_TOTAL),
    ]
    for col_idx, formula, num_fmt, font, fill in tot_cols:
        c = ws_summary.cell(row=curr_row, column=col_idx, value=formula)
        c.font = font
        c.alignment = ALIGN_RIGHT
        c.number_format = num_fmt
        c.fill = fill
        c.border = BORDER_TOTAL_ROW

    curr_row += 2

    # Table 2: Top States in 6-Month Window
    ws_summary.cell(row=curr_row, column=1, value="2. TOP STATES BY SALES TAX CHARGED (6-MONTH WINDOW: MAR 11 – SEP 11, 2026)").font = FONT_SECTION
    ws_summary.row_dimensions[curr_row].height = 20
    curr_row += 1

    state_headers = [
        "State Code", "State Name", "Kelly Approved Status", "Eligible Orders",
        "Tax Refund Requested ($)", "Order Volume ($)", "% of Total Refund"
    ]
    style_table_header(ws_summary, curr_row, state_headers, fill=FILL_TH_SUB)
    header_state_row = curr_row
    curr_row += 1

    state_summary = df_6m.groupby("ship_state_clean").agg(
        order_count=("wb_order_id", "count"),
        tax_sum=("tax", "sum"),
        spend_sum=("order_total", "sum")
    ).reset_index().sort_values(by="tax_sum", ascending=False)

    for _, srow in state_summary.iterrows():
        st_code = srow["ship_state_clean"]
        if not st_code or srow["tax_sum"] <= 0:
            continue
        ws_summary.row_dimensions[curr_row].height = 18
        row_fill = FILL_ZEBRA if curr_row % 2 == 0 else FILL_WHITE
        is_approved = st_code in KELLY_APPROVED_STATES
        st_status = "Approved (Sep 11, 2026)" if is_approved else "Pending State Exemption"

        vals = [
            (st_code, ALIGN_CENTER, None, FONT_TD_BOLD),
            (STATE_NAMES.get(st_code, st_code), ALIGN_LEFT, None, FONT_TD),
            (st_status, ALIGN_CENTER, None, FONT_TD if is_approved else FONT_TD_BOLD),
            (int(srow["order_count"]), ALIGN_RIGHT, FMT_INT, FONT_TD),
            (float(srow["tax_sum"]), ALIGN_RIGHT, FMT_CURRENCY, FONT_TAX_HIGHLIGHT),
            (float(srow["spend_sum"]), ALIGN_RIGHT, FMT_CURRENCY, FONT_TD),
            (f"=E{curr_row}/$F$19", ALIGN_RIGHT, FMT_PCT, FONT_TD),
        ]

        for col_idx, (val, align, num_fmt, font) in enumerate(vals, start=1):
            c = ws_summary.cell(row=curr_row, column=col_idx, value=val)
            c.alignment = align
            c.font = font
            c.border = BORDER_CELL
            c.fill = FILL_TAX_CELL if col_idx == 5 else (FILL_ALERT_CELL if col_idx == 3 and not is_approved else row_fill)
            if num_fmt:
                c.number_format = num_fmt

        curr_row += 1

    # State Table Total Row
    ws_summary.row_dimensions[curr_row].height = 20
    ws_summary.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
    c_st_tot = ws_summary.cell(row=curr_row, column=1, value="TOTAL ALL STATES")
    c_st_tot.font = FONT_TOTAL
    c_st_tot.alignment = Alignment(horizontal="right", vertical="center")
    for col in range(1, 4):
        ws_summary.cell(row=curr_row, column=col).fill = FILL_TOTAL
        ws_summary.cell(row=curr_row, column=col).border = BORDER_TOTAL_ROW

    ws_summary.cell(row=curr_row, column=4, value=f"=SUM(D{header_state_row+1}:D{curr_row-1})").number_format = FMT_INT
    ws_summary.cell(row=curr_row, column=5, value=f"=SUM(E{header_state_row+1}:E{curr_row-1})").number_format = FMT_CURRENCY
    ws_summary.cell(row=curr_row, column=6, value=f"=SUM(F{header_state_row+1}:F{curr_row-1})").number_format = FMT_CURRENCY
    ws_summary.cell(row=curr_row, column=7, value=f"=SUM(G{header_state_row+1}:G{curr_row-1})").number_format = FMT_PCT

    for col in range(4, 8):
        c = ws_summary.cell(row=curr_row, column=col)
        c.font = FONT_TOTAL
        c.alignment = ALIGN_RIGHT
        c.fill = FILL_TAX_CELL if col == 5 else FILL_TOTAL
        c.border = BORDER_TOTAL_ROW

    curr_row += 2

    # Section 3: Notes & Instructions
    ws_summary.cell(row=curr_row, column=1, value="3. NOTES & SUBMISSION DETAILS FOR ZORO ACCOUNTING DEPARTMENT").font = FONT_SECTION
    ws_summary.row_dimensions[curr_row].height = 20
    curr_row += 1

    notes = [
        "1. Exemption Authority: In accordance with email confirmation from Kelly (Zoro Tax Exemption Specialist) on Sep 11, 2026, all 3 customer accounts are coded exempt for 44 states under M&R Market, LLC.",
        "2. Accounting Refund Window: In line with Zoro's 6-month refund policy, this claim specifically presents transactions from March 11, 2026 through September 11, 2026 totaling $50,398.48 across 14,579 WB orders.",
        "3. Tab Layout: Individual tabs are provided for each customer account ('inuhrecomllc', 'salesmercentile', 'salesmercentile1') as well as a consolidated flat claim tab ('All Accounts (6-Month Claim)').",
        "4. Historical Reference: Transactions older than 6 months ($5,556.09) are quarantined into the separate 'Historical (> 6 Months)' tab for state-level claims and are NOT included in the primary 6-month claim total.",
        "5. Remittance: Tax refunds can be credited back to the original payment methods on file or issued as a corporate credit balance to Customer #25017919."
    ]
    for n in notes:
        ws_summary.row_dimensions[curr_row].height = 18
        ws_summary.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=10)
        cn = ws_summary.cell(row=curr_row, column=1, value=n)
        cn.font = FONT_TD
        cn.alignment = ALIGN_LEFT
        curr_row += 1

    ws_summary.views.sheetView[0].showGridLines = True
    auto_fit_columns(ws_summary)

    # =====================================================
    # TAB 2: ALL ACCOUNTS MASTER (6-MONTH CLAIM)
    # =====================================================
    print("  Creating 'All Accounts (6-Month Claim)' tab...")
    ws_all = wb.create_sheet(title="All Accounts (6-Month Claim)")
    create_title_banner(
        ws_all,
        "ALL ACCOUNTS — 6-MONTH SALES TAX REFUND CLAIM ($50,398.48)",
        "Consolidated Order Ledger for Zoro Accounting | March 11, 2026 – September 11, 2026 | M&R Market, LLC",
        max_col=13
    )
    write_order_table(ws_all, df_6m, start_row=4, is_single_account=False)

    # =====================================================
    # TABS 3, 4, 5: INDIVIDUAL ACCOUNT 6-MONTH CLAIMS
    # =====================================================
    for acct_key in ["inhur_llc", "salesmercentile", "salesmercentile1"]:
        meta = ACCOUNTS_META[acct_key]
        df_sub_6m = df_6m[df_6m["gmail_account"] == acct_key].copy()
        sub_tax = df_sub_6m["tax"].sum()
        sub_orders = len(df_sub_6m)
        print(f"  Creating tab for {meta['name']} ({sub_orders} orders, ${sub_tax:,.2f})...")

        ws_sub = wb.create_sheet(title=meta["name"][:31])
        create_title_banner(
            ws_sub,
            f"{meta['name'].upper()} (CUST #{meta['cust_id']}) — 6-MONTH TAX REFUND CLAIM",
            f"Email: {meta['email']} | Business: {meta['entity']} | Total Claim: ${sub_tax:,.2f} ({sub_orders:,} Orders)",
            max_col=10
        )
        write_order_table(ws_sub, df_sub_6m, start_row=4, is_single_account=True)

    # =====================================================
    # TAB 6: HISTORICAL (> 6 MONTHS) REFERENCE
    # =====================================================
    print(f"  Creating 'Historical (> 6 Months)' tab ({len(df_hist)} orders, ${df_hist['tax'].sum():,.2f})...")
    ws_hist = wb.create_sheet(title="Historical (> 6 Months)")
    create_title_banner(
        ws_hist,
        "HISTORICAL TRANSACTIONS PRIOR TO 6-MONTH ACCOUNTING WINDOW",
        "Reference for State Direct Refund Filings | Orders Prior to March 11, 2026 | M&R Market, LLC",
        max_col=13
    )
    write_order_table(ws_hist, df_hist, start_row=4, is_single_account=False)

    wb.save(out_path)
    print(f"Master Workbook saved successfully: {out_path} ({os.path.getsize(out_path)/1024/1024:.2f} MB)")


# ---------------------------------------------------------
# STANDALONE SINGLE-ACCOUNT WORKBOOK BUILDER
# ---------------------------------------------------------
def generate_individual_workbook(df_acct, meta, out_path):
    print(f"Generating Individual Workbook for {meta['name']}: {out_path}...")
    wb = openpyxl.Workbook()
    ws_claim = wb.active
    ws_claim.title = "6-Month Tax Refund Claim"

    df_6m = df_acct[df_acct["is_6m"]].copy()
    df_hist = df_acct[~df_acct["is_6m"]].copy()

    total_6m_tax = df_6m["tax"].sum()
    total_6m_orders = len(df_6m)

    create_title_banner(
        ws_claim,
        f"{meta['name'].upper()} — SALES TAX REFUND REQUEST (${total_6m_tax:,.2f})",
        f"Customer #{meta['cust_id']} | Email: {meta['email']} | Business: {meta['entity']} | Window: Mar 11 – Sep 11, 2026",
        max_col=10
    )
    write_order_table(ws_claim, df_6m, start_row=4, is_single_account=True)

    if len(df_hist) > 0:
        ws_hist = wb.create_sheet(title="Historical (> 6 Months)")
        create_title_banner(
            ws_hist,
            f"{meta['name'].upper()} — HISTORICAL TRANSACTIONS (> 6 MONTHS)",
            f"Customer #{meta['cust_id']} | Orders Prior to March 11, 2026 | Total Tax: ${df_hist['tax'].sum():,.2f}",
            max_col=10
        )
        write_order_table(ws_hist, df_hist, start_row=4, is_single_account=True)

    wb.save(out_path)
    print(f"Saved {meta['name']} workbook: {out_path} ({os.path.getsize(out_path)/1024/1024:.2f} MB)")


# ---------------------------------------------------------
# MAIN EXECUTION
# ---------------------------------------------------------
def main():
    db_path = "Automations/data/automations_unified.db"
    out_dir = "Automations/data"
    os.makedirs(out_dir, exist_ok=True)

    print(f"Loading Zoro tax data from {db_path}...")
    df = load_tax_data(db_path)
    print(f"Loaded {len(df):,} total taxed orders across 3 accounts.")
    print(f"  6-Month Eligible (>= {CUTOFF_6M}): {df['is_6m'].sum():,} orders | Tax: ${df[df['is_6m']]['tax'].sum():,.2f}")
    print(f"  Historical (> 6 Months): {(~df['is_6m']).sum():,} orders | Tax: ${df[~df['is_6m']]['tax'].sum():,.2f}")

    # 1. Master Consolidated Workbook
    master_path = os.path.join(out_dir, "Zoro_Sales_Tax_Refund_Master_MR_Market_LLC.xlsx")
    generate_master_workbook(df, master_path)

    # 2. Standalone Individual Workbooks
    for acct_key, meta in ACCOUNTS_META.items():
        df_acct = df[df["gmail_account"] == acct_key].copy()
        ind_path = os.path.join(out_dir, f"Zoro_Sales_Tax_Refund_{meta['name']}_CUST{meta['cust_id']}.xlsx")
        generate_individual_workbook(df_acct, meta, ind_path)

    print("\nAll workbooks generated successfully!")

if __name__ == "__main__":
    main()
