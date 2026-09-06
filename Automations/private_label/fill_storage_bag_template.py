#!/usr/bin/env python3
"""fill_storage_bag_template.py — fill Amazon's STORAGE_BAG category template
with HOANK reclosable zip-seal bag listings built straight from
combined_reports.db (288 Zoro Select listings).

Follows the conventions proven by fill_shipping_box_template.py:
new-format template (labels row 4, attributes row 5, data from row 8),
DISPLAY dropdown values, SET_NAME variation theme, parents stripped of
offer/id fields. UPC and images intentionally left blank (added later);
product id type = GTIN Exempt for buyable rows.

Category: Industrial & Scientific > Packaging & Shipping Supplies >
Poly & Plastic Packaging Bags  -> item_type_keyword
"professional-business-poly-packaging-bags" (from the template's own list).

    python3 fill_storage_bag_template.py                 # all 288
    python3 fill_storage_bag_template.py --limit 5       # test run
"""
import argparse
import os
import re
import sqlite3
from collections import defaultdict
from fractions import Fraction

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ZORO_DB = os.path.join(HERE, "..", "TSV Files", "combined_reports.db")
TEMPLATE = os.path.expanduser("~/Downloads/STORAGE_BAG.xlsm")
OUT_DEFAULT = os.path.expanduser("~/Downloads/STORAGE_BAG_HOANK_filled.xlsm")
ATTR_ROW, DATA_ROW = 5, 8

MK = "[marketplace_id=ATVPDKIKX0DER]"
LT = "[language_tag=en_US]"
COO_MAP = {"US": "United States", "CN": "China", "VN": "Vietnam",
           "TW": "Taiwan", "PH": "Philippines"}   # ZZ/unknown -> blank
RELEASE = "2026-08-01"
SHIP_TEMPLATE = "Main PL Tamplate 18.99"          # Amazon-side name (sic)
MARKUP = 1.11                                     # same base curve as boxes


def num(s):
    s = s.strip()
    if "/" in s:
        parts = s.split()
        whole = float(parts[0]) if len(parts) > 1 else 0.0
        return whole + float(Fraction(parts[-1]))
    return float(s)


def fmt(v):
    return "%g" % round(v, 2)


def fetch_rows():
    con = sqlite3.connect(f"file:{ZORO_DB}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    return con.execute("""
        SELECT zoro_no, short_description, long_description, sales_price,
               upc_code, country_of_origin, ship_pack_length, ship_pack_width,
               ship_pack_height, ship_pack_weight
        FROM zoro_data
        WHERE zoro_brand_name='Zoro Select' AND availabilityStatus='In Stock'
          AND is_dropship='NO'
          AND (lead_time LIKE '%ships in 1 day%' OR lead_time LIKE '%Ships in 1 day%')
          AND short_description LIKE '%Reclosable Bag%'
        """).fetchall()


def parse_bag(r):
    ld = r["long_description"] or ""
    sd = r["short_description"] or ""

    def grab(pat, src=None):
        m = re.search(pat, src if src is not None else ld)
        return m.group(1).strip() if m else None

    # specs usually live in long_description; some rows carry them in the
    # short_description instead ("18 in L, 20 in W, 4 mil Thick, ...")
    w = grab(r"Width ([\d\s/\.]+) in") or grab(r"([\d\s/\.]+) in W\b", sd)
    l = grab(r"Length ([\d\s/\.]+) in") or grab(r"([\d\s/\.]+) in L\b", sd)
    mil = grab(r"Thickness ([\d\.]+) mil") or grab(r"([\d\.]+) mil Thick", sd)
    pack = grab(r"(?:Package|Pack) (?:Quantity|Qty\.?) ([\d,]+)") or \
        grab(r"(\d+)\s*PK", sd)
    if not (w and l and pack):
        return None
    feat = grab(r"Bag Features ([^,]+)") or \
        ("White Block" if "White Block" in ld else "Standard")
    fda = "FDA Compliant" in ld and "Not FDA Compliant" not in ld
    color = grab(r",\s*(Clear|Pink|Blue|Black|White|Green|Red|Amber|Silver)\s*$", sd) \
        or "Clear"
    hang = bool(re.search(r"Hang Hole Included", ld))
    return dict(W=num(w), L=num(l), mil=mil or "2", pack=int(pack.replace(",", "")),
                feat=feat, fda=fda, color=color, hang=hang)


def feature_phrase(d):
    f = d["feat"]
    if f == "White Block":
        return "with Writable White Block"
    if f == "Anti-Static":
        return "Anti-Static"
    if "VCI" in f:
        return "Corrosion-Inhibiting (VCI)"
    return ""


def build_title(d):
    size = f'{fmt(d["W"])}" x {fmt(d["L"])}"'
    fp = feature_phrase(d)
    bits = [f"HOANK Reclosable Zip Bags {size}", f'{d["pack"]} Pack',
            f'{d["mil"]} Mil {d["color"]} Plastic Zipper Seal Bags']
    if fp:
        bits.append(fp)
    bits.append("Resealable Poly Bags for Jewelry, Beads, Small Parts, Storage")
    t = ", ".join(bits)
    return t[:199]


def build_bullets(d):
    size = f'{fmt(d["W"])}"W x {fmt(d["L"])}"L'
    b = [
        f'BULK PACK OF {d["pack"]}: HOANK reclosable bags come in a case of '
        f'{d["pack"]}, a long-lasting supply for organizing, shipping, and '
        f'storing small items at home or in the shop.',
        f'SIZE: Each bag measures {size}, the right fit for jewelry, beads, '
        f'coins, hardware, pills, crafts, and other small parts.',
        f'DURABLE {d["mil"]} MIL FILM: Made of {d["mil"]} mil low density '
        f'polyethylene (LDPE) that resists tears and punctures while staying '
        f'flexible and crystal {d["color"].lower()}.',
        'SECURE ZIP SEAL: The press-to-close zipper opens and reseals '
        'hundreds of times, keeping contents in and dust and moisture out.',
    ]
    extra = []
    if d["feat"] == "White Block":
        extra.append("WRITABLE WHITE BLOCK: A built-in white panel takes pen "
                     "or marker, so every bag can be labeled and found fast.")
    elif d["feat"] == "Anti-Static":
        extra.append("ANTI-STATIC PROTECTION: Dissipates static charge to "
                     "keep sensitive electronic components and parts safe.")
    elif "VCI" in d["feat"]:
        extra.append("VCI CORROSION PROTECTION: Vapor corrosion inhibitors "
                     "guard metal parts against rust during storage.")
    elif d["fda"]:
        extra.append("FDA COMPLIANT FILM: Safe for direct food contact - "
                     "portion snacks, spices, and leftovers with confidence.")
    else:
        extra.append("VERSATILE: Ideal for retail, e-commerce, inventory, "
                     "travel, crafts, and everyday household organization.")
    return b + extra


def build_description(d):
    fp = feature_phrase(d)
    return (
        f'HOANK {fmt(d["W"])}" x {fmt(d["L"])}" reclosable zip seal bags, '
        f'{d["pack"]} per case. Made from {d["mil"]} mil low density '
        f'polyethylene (LDPE), these {d["color"].lower()} poly bags combine '
        f'see-through clarity with dependable strength{", " + fp.lower() if fp else ""}. '
        f'The press-to-close zipper seals securely and reopens smoothly, '
        f'protecting contents from dust, moisture, and spills. Stock up once '
        f'and keep jewelry, beads, coins, hardware, samples, crafts, and '
        f'small parts sorted, protected, and easy to find - at home, in the '
        f'stockroom, or at the point of sale.'
    )


def build_row(r, d, sku, price, qty):
    coo = COO_MAP.get((r["country_of_origin"] or "").strip(), "")
    v = {
        "contribution_sku#1.value": sku,
        "product_type#1.value": "STORAGE_BAG",
        "::record_action": "Create or Replace (Full Update)",
        f"item_name{MK}{LT}#1.value": build_title(d),
        f"brand{MK}{LT}#1.value": "HOANK",
        f"manufacturer{MK}{LT}#1.value": "HOANK",
        "amzn1.volt.ca.product_id_type": "GTIN Exempt",
        f"item_type_keyword{MK}#1.value": "professional-business-poly-packaging-bags",
        f"model_number{MK}#1.value": f'HO-K{r["zoro_no"]}-{d["pack"]}',
        f"model_name{MK}{LT}#1.value":
            f'{fmt(d["W"])}x{fmt(d["L"])} {d["mil"]} Mil {d["pack"]} Pack',
        f"is_fulfilled_from_origin_item{MK}#1.value": "Yes",
        f"street_date{MK}#1.value": RELEASE,
        f"merchant_release_date{MK}#1.value": RELEASE,
        "fulfillment_availability#1.fulfillment_channel_code":
            "Fulfillment by Merchant (Default)",
        "fulfillment_availability#1.quantity": qty,
        f"list_price{MK}#1.value": price,
        "purchasable_offer[marketplace_id=ATVPDKIKX0DER][audience=ALL]"
        "#1.our_price#1.schedule#1.value_with_tax": price,
        f"merchant_shipping_group{MK}#1.value": SHIP_TEMPLATE,
        f"product_description{MK}{LT}#1.value": build_description(d),
        f"supplier_declared_dg_hz_regulation{MK}#1.value": "Not Applicable",
        f"condition_type{MK}#1.value": "New",
        f"color{MK}{LT}#1.value": d["color"],
        f"size{MK}{LT}#1.value":
            f'{fmt(d["W"])}" x {fmt(d["L"])}" ({d["pack"]} Pack)',
        f"material{MK}{LT}#1.value": "Low Density Polyethylene (LDPE)",
        f"closure{MK}#1.type{LT}#1.value": "Zipper",
        f"number_of_items{MK}#1.value": d["pack"],
        f"item_package_quantity{MK}#1.value": 1,
        f"unit_count{MK}#1.value": d["pack"],
        f"unit_count{MK}#1.type{LT}.value": "Count",
        f"number_of_packs{MK}#1.value": d["pack"],
        f"item_length_width_height{MK}#1.length.value": fmt(d["L"]),
        f"item_length_width_height{MK}#1.length.unit": "Inches",
        f"item_length_width_height{MK}#1.width.value": fmt(d["W"]),
        f"item_length_width_height{MK}#1.width.unit": "Inches",
    }
    if coo:
        v[f"country_of_origin{MK}#1.value"] = coo
    sf = feature_phrase(d)
    if sf:
        v[f"special_feature{MK}{LT}#1.value"] = sf
    for i, u in enumerate(["Jewelry", "Beads", "Small Parts",
                           "Coins", "Crafts"], 1):
        v[f"recommended_uses_for_product{MK}{LT}#{i}.value"] = u
    if r["ship_pack_weight"]:
        v[f"item_weight{MK}#1.value"] = r["ship_pack_weight"]
        v[f"item_weight{MK}#1.unit"] = "Pounds"
        v[f"item_package_weight{MK}#1.value"] = r["ship_pack_weight"]
        v[f"item_package_weight{MK}#1.unit"] = "Pounds"
    if r["ship_pack_length"]:
        v[f"item_package_dimensions{MK}#1.length.value"] = r["ship_pack_length"]
        v[f"item_package_dimensions{MK}#1.length.unit"] = "Inches"
        v[f"item_package_dimensions{MK}#1.width.value"] = r["ship_pack_width"]
        v[f"item_package_dimensions{MK}#1.width.unit"] = "Inches"
        v[f"item_package_dimensions{MK}#1.height.value"] = r["ship_pack_height"]
        v[f"item_package_dimensions{MK}#1.height.unit"] = "Inches"
    for i, b in enumerate(build_bullets(d), 1):
        v[f"bullet_point{MK}{LT}#{i}.value"] = b
    for i, kw in enumerate(["small plastic bags", "resealable bags",
                            "zip lock bags bulk", "clear poly bags",
                            "jewelry bags"], 1):
        v[f"generic_keyword{MK}{LT}#{i}.value"] = kw
    return v


PARENT_DROP = ("purchasable_offer", "fulfillment_availability", "list_price",
               "merchant_shipping_group", "product_id", "max_order_quantity",
               "number_of_items", "unit_count", "number_of_packs", "set_name",
               "item_package_quantity", "model_number", "model_name",
               "item_weight", "item_package_weight", "item_package_dimensions",
               "closure", "special_feature", "country_of_origin",
               "item_length_width_height", "recommended_uses")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=None,
                    help="test run: only this many size families")
    ap.add_argument("--qty", type=int, default=1000)
    ap.add_argument("--out", default=OUT_DEFAULT)
    a = ap.parse_args()

    rows = fetch_rows()
    parsed = []
    for r in rows:
        d = parse_bag(r)
        if d:
            parsed.append((r, d))
        else:
            print(f"  !! skipped (unparseable): {r['zoro_no']}  "
                  f"{(r['short_description'] or '')[:70]}")
    print(f"{len(parsed)} of {len(rows)} bags parsed")

    fams = defaultdict(list)
    for r, d in parsed:
        fams[(d["W"], d["L"])].append((r, d))
    keys = sorted(fams, key=lambda k: (k[0], k[1]))
    if a.limit:
        keys = keys[:a.limit]
        print(f"--limit: {a.limit} families "
              f"({sum(len(fams[k]) for k in keys)} bags)")

    wb = openpyxl.load_workbook(TEMPLATE, keep_vba=True)
    ws = wb["Template"]
    attr_col = {}
    for ci, cell in enumerate(ws[ATTR_ROW], 1):
        if cell.value:
            attr_col[str(cell.value).strip()] = ci

    unknown, ri, n_parent, n_child, n_single = set(), 0, 0, 0, 0
    for key in keys:
        members = sorted(fams[key], key=lambda rd:
                         (float(rd[1]["mil"]), rd[1]["feat"], rd[1]["pack"]))
        W, L = key
        is_family = len(members) > 1
        parent_sku = f"HOA-Kam-P-BAG-{fmt(W)}x{fmt(L)}"
        if is_family:
            r0, d0 = members[0]
            pv = build_row(r0, d0, parent_sku, 0, 0)
            pv[f"item_name{MK}{LT}#1.value"] = (
                f'HOANK Reclosable Zip Bags {fmt(W)}" x {fmt(L)}", Resealable '
                f'Zipper Seal Poly Bags for Jewelry, Beads, Small Parts, Storage')
            pv[f"parentage_level{MK}#1.value"] = "Parent"
            pv["variation_theme#1.name"] = "SET_NAME"
            pv["amzn1.volt.ca.product_id_type"] = None
            for k in [k for k in list(pv) if any(t in k for t in PARENT_DROP)]:
                del pv[k]
            for attr, ci in attr_col.items():
                ws.cell(row=DATA_ROW + ri, column=ci).value = pv.get(attr)
            unknown |= {x for x in pv if x not in attr_col}
            ri += 1
            n_parent += 1

        seen = {}
        for r, d in members:
            sku = f'HOA-Kam-{r["zoro_no"]}-{d["pack"]}'
            price = round(float(r["sales_price"] or 0) * MARKUP, 2)
            v = build_row(r, d, sku, price, a.qty)
            if is_family:
                v[f"parentage_level{MK}#1.value"] = "Child"
                v[f"child_parent_sku_relationship{MK}#1.parent_sku"] = parent_sku
                v["variation_theme#1.name"] = "SET_NAME"
                name = f'{d["mil"]} Mil'
                fp = feature_phrase(d)
                if fp:
                    name += f" - {fp}"
                name += f' ({d["pack"]} Pack)'
                n = seen.get(name, 0) + 1
                seen[name] = n
                if n > 1:
                    name = f"{name} (Style {n})"
                v[f"set_name{MK}{LT}#1.value"] = name
                n_child += 1
            else:
                n_single += 1
            for attr, ci in attr_col.items():
                ws.cell(row=DATA_ROW + ri, column=ci).value = v.get(attr)
            unknown |= {x for x in v if x not in attr_col}
            ri += 1

    if unknown:
        print("⚠️ attributes not found in template (skipped):")
        for u in sorted(unknown):
            print("   ", u)
    wb.save(a.out)
    print(f"\nwrote {ri} rows ({n_parent} parents, {n_child} children, "
          f"{n_single} standalone) -> {a.out}")


if __name__ == "__main__":
    main()
