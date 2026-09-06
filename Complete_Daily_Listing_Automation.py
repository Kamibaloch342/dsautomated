import pandas as pd
import os
import re
import sqlite3
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
from ttkthemes import ThemedTk
from pathlib import Path
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.worksheet.table import Table, TableStyleInfo
import threading

# --- 1. CORE PROCESSING LOGIC ---

def process_all_listings_report(file_path):
    cols_to_keep = ['item-name', 'seller-sku', 'price', 'quantity', 'open-date', 'asin1', 'status']
    df = pd.read_csv(file_path, sep='\t', usecols=cols_to_keep, on_bad_lines='warn', encoding='utf-8')
    df['open-date'] = pd.to_datetime(df['open-date'], errors='coerce').dt.strftime('%d %B %Y')
    df.rename(columns={
        'item-name': 'Product_Title', 'seller-sku': 'SKU', 'price': 'Selling_Price',
        'quantity': 'Qty', 'open-date': 'Open_Date', 'asin1': 'ASIN', 'status': 'Status'
    }, inplace=True)
    return df

def parse_and_normalize_skus(df):
    def parse_sku_details(sku):
        if not isinstance(sku, str):
            return None, 1
        
        zoro_code_match = re.search(r'(G\d{7,}|ZOG\d{7,})', sku)
        zoro_code = zoro_code_match.group(1) if zoro_code_match else None
        
        if zoro_code and zoro_code.startswith('ZOG'):
            zoro_code = 'G' + zoro_code[3:]
            
        pack_numbers = re.findall(r'-(\d+)', sku)
        if pack_numbers:
            pack = int(pack_numbers[-1])
            if pack == 0: pack = 1
        else:
            pack = 1
        
        return zoro_code, pack

    parsed_data = df['SKU'].apply(parse_sku_details)
    df[['Zoro_Code', 'Pack']] = pd.DataFrame(parsed_data.tolist(), index=df.index)
    df['Normalized_SKU'] = df.apply(
        lambda row: f"{row['Zoro_Code']}-{row['Pack']}" if pd.notna(row['Zoro_Code']) else None,
        axis=1
    )
    return df

def create_zoro_database(tsv_folder_path, db_path, logger):
    if db_path.exists():
        logger("Zoro database already exists. Skipping creation.")
        return
    logger("Creating Zoro database from TSV files. This may take a moment...")
    tsv_files = list(Path(tsv_folder_path).glob('*.tsv'))
    if not tsv_files: raise FileNotFoundError(f"No .tsv files found in {tsv_folder_path}")

    df_first = pd.read_csv(tsv_files[0], sep='\t', on_bad_lines='warn', low_memory=False)
    headers = df_first.columns.tolist()
    def read_other_files():
        yield df_first
        for file in tsv_files[1:]:
            yield pd.read_csv(file, sep='\t', header=None, names=headers, on_bad_lines='warn', low_memory=False)
    combined_df = pd.concat(read_other_files(), ignore_index=True)
    with sqlite3.connect(db_path) as conn:
        combined_df.to_sql("zoro_data", conn, if_exists="replace", index=False)
    logger("-> Zoro database created successfully.")

def process_order_reports(folder_path, logger):
    logger("Processing order reports to calculate sales data...")
    all_files = list(Path(folder_path).glob("*.txt"))
    if not all_files:
        logger("-> No order reports found. Sales data will be zero.")
        return pd.DataFrame(columns=['ASIN', 'Total_Units_Sold', 'SKU'])
        
    df_list = [pd.read_csv(f, sep='\t', on_bad_lines='warn', encoding='utf-8', low_memory=False) for f in all_files]
    if not df_list: return pd.DataFrame(columns=['ASIN', 'Total_Units_Sold', 'SKU'])
        
    orders_df = pd.concat(df_list, ignore_index=True)
    shipped = orders_df[orders_df['order-status'] == 'Shipped'].copy()
    shipped['quantity'] = pd.to_numeric(shipped['quantity'], errors='coerce').fillna(0)
    
    sales_by_asin = shipped.groupby('asin').agg(
        Total_Units_Sold=('quantity', 'sum'),
        SKU=('sku', 'last') 
    ).reset_index()
    sales_by_asin.rename(columns={'asin': 'ASIN'}, inplace=True)
    
    logger("-> Order reports processed.")
    return sales_by_asin

def analyze_duplicates(df):
    skus_to_delete = []
    duplicate_asins = df[df.duplicated(subset=['ASIN'], keep=False)]['ASIN'].unique()
    for asin in duplicate_asins:
        group = df[df['ASIN'] == asin].copy()
        skus_to_delete.extend(group[group['Status'] == 'Inactive']['SKU'].tolist())
        active_group = group[group['Status'] == 'Active']
        if len(active_group) <= 1: continue
        sorted_by_sales = active_group.sort_values(by='Total_Units_Sold', ascending=False)
        highest_sales = sorted_by_sales['Total_Units_Sold'].iloc[0]
        skus_to_delete.extend(sorted_by_sales[sorted_by_sales['Total_Units_Sold'] < highest_sales]['SKU'].tolist())
        top_sellers = sorted_by_sales[sorted_by_sales['Total_Units_Sold'] == highest_sales]
        if len(top_sellers) > 1:
            kam_skus = top_sellers[top_sellers['SKU'].str.startswith('Kam', na=False)]
            sku_to_keep = kam_skus.iloc[0]['SKU'] if not kam_skus.empty else top_sellers.iloc[0]['SKU']
            skus_to_delete.extend(top_sellers[top_sellers['SKU'] != sku_to_keep]['SKU'].tolist())
    return list(set(skus_to_delete))

def perform_profit_analysis(df):
    df['Selling_Price'] = pd.to_numeric(df['Selling_Price'], errors='coerce')
    df['Zoro_Price'] = pd.to_numeric(df['sales_price'], errors='coerce')
    df['Cost_Price'] = df['Zoro_Price'] * df['Pack']
    df['Referral_Fee'] = 0.15 * df['Selling_Price']
    df['Shipping_Fee'] = df['Cost_Price'].apply(lambda x: 5 if pd.notna(x) and x < 50 else 0)
    df['Total_Cost'] = df['Cost_Price'] + df['Referral_Fee'] + df['Shipping_Fee']
    df['Profit'] = df['Selling_Price'] - df['Total_Cost']
    df['Profit_Margin'] = (df['Profit'] / df['Selling_Price']).replace([float('inf'), -float('inf')], 0)
    # --- MODIFIED FOR 15% MARGIN ---
    df['Suggested_Price_15_Margin'] = pd.NA
    low_profit_mask = df['Profit_Margin'] < 0.15
    df.loc[low_profit_mask, 'Suggested_Price_15_Margin'] = ((df.loc[low_profit_mask, 'Cost_Price'] + df.loc[low_profit_mask, 'Shipping_Fee']) / 0.70).round(2)
    return df

def format_excel_sheet(worksheet, df):
    header_font = Font(bold=True, size=12, name='Calibri', color="FFFFFF")
    cell_font = Font(size=10, name='Calibri')
    alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    formats = {
        'Selling_Price': '$#,##0.00', 'Zoro_Price': '$#,##0.00', 'Cost_Price': '$#,##0.00',
        'Profit': '$#,##0.00', 'Suggested_Price_15_Margin': '$#,##0.00', 'Referral_Fee': '$#,##0.00',
        'Shipping_Fee': '$#,##0.00', 'Total_Cost': '$#,##0.00',
        'Profit_Margin': '0.00%'
    }

    for row_idx, row in enumerate(worksheet.iter_rows(), 1):
        for cell in row:
            if row_idx == 1:
                cell.font = header_font
                cell.fill = PatternFill(start_color="A93226", end_color="A93226", fill_type="solid")
            else:
                cell.font = cell_font
            cell.alignment = alignment
            cell.border = thin_border
            col_title = worksheet.cell(row=1, column=cell.column).value
            if col_title in formats:
                cell.number_format = formats[col_title]

    for col_cells in worksheet.columns:
        col_letter = col_cells[0].column_letter
        col_title = worksheet.cell(row=1, column=col_cells[0].column).value
        if col_title == 'Product_Title':
            worksheet.column_dimensions[col_letter].width = 60
        else:
            max_length = max((len(str(cell.value)) for cell in col_cells if cell.value), default=10)
            worksheet.column_dimensions[col_letter].width = max_length + 2

    table = Table(displayName=f"Table_{worksheet.title.replace(' ', '_').replace('%', '')}", ref=worksheet.dimensions)
    style = TableStyleInfo(name="TableStyleMedium4", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table.tableStyleInfo = style
    worksheet.add_table(table)
    
    if "Profit_Margin" in df.columns:
        rule = ColorScaleRule(start_type='num', start_value=0.0, start_color='F8696B', mid_type='num', mid_value=0.15, mid_color='FFEB84', end_type='num', end_value=0.25, end_color='63BE7B')
        col_idx = df.columns.get_loc("Profit_Margin") + 1
        worksheet.conditional_formatting.add(f"{chr(64+col_idx)}2:{chr(64+col_idx)}{worksheet.max_row}", rule)
    if "Total_Units_Sold" in df.columns:
        rule = ColorScaleRule(start_type='min', start_color='FFF8E5', end_type='percentile', end_value=95, end_color='63BE7B')
        col_idx = df.columns.get_loc("Total_Units_Sold") + 1
        worksheet.conditional_formatting.add(f"{chr(64+col_idx)}2:{chr(64+col_idx)}{worksheet.max_row}", rule)
    if "Comments" in df.columns:
        col_letter = chr(64 + df.columns.get_loc("Comments") + 1)
        range_str = f"{col_letter}2:{col_letter}{worksheet.max_row}"
        rules_map = {'🔴': "F8696B", '🟢': "63BE7B", '💰': "FFEB84", '🗑️': "A93226"}
        for key, color in rules_map.items():
            worksheet.conditional_formatting.add(range_str, CellIsRule(operator='containsText', formula=[f'"{key}"'], fill=PatternFill(start_color=color, end_color=color, fill_type="solid")))
    worksheet.sheet_view.showGridLines = False

# --- 2. GUI APPLICATION CLASS ---
class App(ThemedTk):
    def __init__(self):
        super().__init__()
        self.set_theme("arc")
        self.title("Super Amazon Analyzer")
        self.geometry("800x600")

        self.paths = {"listings": "", "orders": "", "zoro": ""}

        main_frame = ttk.Frame(self, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)

        self.files_frame = ttk.LabelFrame(main_frame, text="1. Select Files & Folders", padding="10")
        self.files_frame.pack(fill=tk.X, expand=True, pady=5)
        self.files_frame.grid_columnconfigure(1, weight=1)

        self.listings_label = self.create_path_selector(self.files_frame, "All Listings Report:", 0, "listings")
        self.orders_label = self.create_path_selector(self.files_frame, "Order Reports Folder:", 1, "orders")
        self.zoro_label = self.create_path_selector(self.files_frame, "Zoro Data Folder:", 2, "zoro")
        
        self.run_button = ttk.Button(main_frame, text="Run Analysis", state="disabled", command=self.start_analysis_thread)
        self.run_button.pack(pady=10, fill=tk.X, expand=True)

        log_frame = ttk.LabelFrame(main_frame, text="Progress Log", padding="10")
        log_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        self.log_area = scrolledtext.ScrolledText(log_frame, wrap=tk.WORD, height=15, state="disabled")
        self.log_area.pack(fill=tk.BOTH, expand=True)

    def create_path_selector(self, parent, text, row, path_key):
        ttk.Label(parent, text=text).grid(row=row, column=0, sticky="w", padx=5, pady=5)
        label = ttk.Label(parent, text="Not selected", anchor="w")
        label.grid(row=row, column=1, sticky="ew", padx=5)
        ttk.Button(parent, text="Browse...", command=lambda: self.browse(path_key)).grid(row=row, column=2, sticky="e", padx=5)
        return label
    
    def browse(self, key):
        label_text = self.files_frame.grid_slaves(row=list(self.paths.keys()).index(key), column=0)[0].cget("text")
        if "Folder" in label_text:
            path = filedialog.askdirectory()
        else:
            path = filedialog.askopenfilename()
        
        if path:
            self.paths[key] = path
            label_widget = getattr(self, f"{key}_label")
            label_widget.config(text=os.path.basename(path))
            if all(self.paths.values()):
                self.run_button.config(state="normal")
    
    def log(self, message):
        self.log_area.config(state="normal")
        self.log_area.insert(tk.END, message + "\n")
        self.log_area.config(state="disabled")
        self.log_area.see(tk.END)
        self.update_idletasks()
    
    def start_analysis_thread(self):
        self.run_button.config(state="disabled")
        self.log("Starting analysis... The application may be unresponsive for a moment.")
        thread = threading.Thread(target=self.run_analysis)
        thread.daemon = True
        thread.start()
        
    def run_analysis(self):
        try:
            # --- ADDED ---: Blacklists for ASINs and Brands
            BLACKLISTED_ASINS = {
                'B00NQWLGJI', 'B01LQKA250', 'B01LWJYG34', 'B000EDTS1A', 'B004JVV3KC',
                'B00GYW5HH4', 'B00NL48TFK', 'B00NL48UPY', 'B00NL4275S', 'B00NL4DI2E',
                'B00NL4AAMA', 'B014RDA4O2', 'B07P6W6JLF', 'B007IABE5K', 'B0106A4TSC',
                'B088D2MR24', 'B087WGY98W', 'B07BBQYVPT', 'B00FZGNUBK', 'B0036ZBAV4'
            }
            # Keywords are checked case-insensitively
            BLACKLISTED_BRANDS = ['turtle wax', 'trim-lok', 'trim lok']

            self.log("Processing All Listings Report...")
            listings_df = process_all_listings_report(self.paths["listings"])
            listings_df = parse_and_normalize_skus(listings_df)

            sales_df = process_order_reports(self.paths["orders"], self.log)

            db_path = Path(self.paths["zoro"]) / "zoro_combined.db"
            create_zoro_database(self.paths["zoro"], db_path, self.log)
            self.log("Loading data from Zoro database...")
            with sqlite3.connect(db_path) as conn:
                zoro_df = pd.read_sql_query("SELECT zoro_no, sales_price, availabilityStatus, lead_time FROM zoro_data", conn)
                zoro_df['zoro_no'] = zoro_df['zoro_no'].astype(str)
            
            self.log("Merging all data sources...")
            main_df = pd.merge(listings_df, sales_df, on=['ASIN', 'SKU'], how='left')
            main_df['Total_Units_Sold'] = main_df['Total_Units_Sold'].fillna(0).astype(int)
            main_df = pd.merge(main_df, zoro_df, left_on='Zoro_Code', right_on='zoro_no', how='left')

            self.log("Performing profit analysis...")
            main_df = perform_profit_analysis(main_df)
            
            self.log("Analyzing for duplicates, blacklisted items, and missing data...")
            
            duplicate_skus = analyze_duplicates(main_df.copy())
            no_data_skus = main_df[main_df['sales_price'].isna()]['SKU'].tolist()
            
            brand_mask = main_df['Product_Title'].str.lower().str.contains('|'.join(BLACKLISTED_BRANDS), na=False)
            asin_mask = main_df['ASIN'].isin(BLACKLISTED_ASINS)
            blacklisted_skus = main_df[brand_mask | asin_mask]['SKU'].tolist()
            
            all_skus_to_delete = list(set(duplicate_skus + no_data_skus + blacklisted_skus))
            
            self.log("Generating action comments...")
            main_df['Comments'] = '✅ Good to Go' 

            main_df.loc[main_df['SKU'].isin(no_data_skus), 'Comments'] = '🗑️ Delete (No Zoro Data)'
            main_df.loc[main_df['SKU'].isin(duplicate_skus), 'Comments'] = '🗑️ Delete (Duplicate)'
            main_df.loc[main_df['SKU'].isin(blacklisted_skus), 'Comments'] = '🗑️ Delete (Blacklisted ASIN/Brand)'

            safe_to_modify_mask = ~main_df['SKU'].isin(all_skus_to_delete)

            main_df.loc[safe_to_modify_mask & (main_df['Status'] == 'Active') & (main_df['availabilityStatus'] != 'In Stock'), 'Comments'] = '🔴 Make Inactive (OOS on Zoro)'
            
            main_df.loc[safe_to_modify_mask & (main_df['Status'] == 'Inactive') & (main_df['availabilityStatus'] == 'In Stock'), 'Comments'] = '🟢 Make Active (In Stock on Zoro)'
            
            # --- MODIFIED FOR 15% MARGIN ---
            main_df.loc[(main_df['Comments'].isin(['✅ Good to Go', '🟢 Make Active (In Stock on Zoro)'])) & (main_df['Profit_Margin'] < 0.15), 'Comments'] = '💰 Pricing Review: Low Profit'
            main_df.loc[(main_df['Comments'].isin(['✅ Good to Go', '🟢 Make Active (In Stock on Zoro)'])) & (main_df['Profit_Margin'] > 0.30), 'Comments'] = '💰 Pricing Review: High Profit'


            self.log("Preparing 'Sold but Deleted' relisting data...")
            sold_but_deleted_df = sales_df[sales_df['ASIN'].isin(set(sales_df['ASIN']) - set(listings_df['ASIN']))].copy()
            if not sold_but_deleted_df.empty:
                sold_but_deleted_df = parse_and_normalize_skus(sold_but_deleted_df)
                sold_but_deleted_df = pd.merge(sold_but_deleted_df, zoro_df, left_on='Zoro_Code', right_on='zoro_no', how='left')
                sold_but_deleted_df['sales_price'] = pd.to_numeric(sold_but_deleted_df['sales_price'], errors='coerce')
                sold_but_deleted_df['Cost_Price'] = sold_but_deleted_df['sales_price'] * sold_but_deleted_df['Pack']
                sold_but_deleted_df['Shipping_Fee'] = sold_but_deleted_df['Cost_Price'].apply(lambda x: 5 if pd.notna(x) and x < 50 else 0)
                # --- MODIFIED FOR 15% MARGIN ---
                sold_but_deleted_df['Suggested_Price_15_Margin'] = ((sold_but_deleted_df['Cost_Price'] + sold_but_deleted_df['Shipping_Fee']) / 0.70).round(2)

            self.log("Sorting and preparing final sheets...")
            main_df.sort_values(by='Comments', inplace=True)
            
            sheets_to_create = {
                "Master_Report": main_df,
                "Listings_To_Delete": main_df[main_df['SKU'].isin(all_skus_to_delete)],
                "Make_Inactive": main_df[main_df['Comments'] == '🔴 Make Inactive (OOS on Zoro)'],
                "Make_Active": main_df[main_df['Comments'] == '🟢 Make Active (In Stock on Zoro)'],
                "Review_Low_Profit": main_df[main_df['Comments'] == '💰 Pricing Review: Low Profit'].sort_values('Profit_Margin', ascending=True),
                "Review_High_Profit": main_df[main_df['Comments'] == '💰 Pricing Review: High Profit'].sort_values('Profit_Margin', ascending=False),
                "SKU_Normalization": main_df[['SKU', 'Zoro_Code', 'Pack', 'Normalized_SKU']].dropna(subset=['Zoro_Code']),
                "Sold_but_Deleted_Relist": sold_but_deleted_df
            }

            save_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
            if not save_path:
                self.log("Save cancelled. Analysis complete.")
                messagebox.showinfo("Cancelled", "Save operation was cancelled.")
                self.run_button.config(state="normal")
                return

            self.log(f"Writing final report to {os.path.basename(save_path)}...")
            with pd.ExcelWriter(save_path, engine='openpyxl') as writer:
                for sheet_name, df_sheet in sheets_to_create.items():
                    if not df_sheet.empty:
                        if sheet_name == 'Sold_but_Deleted_Relist':
                            # --- MODIFIED FOR 15% MARGIN ---
                            cols_to_show = ['ASIN', 'SKU', 'Total_Units_Sold', 'Zoro_Code', 'Pack', 'Zoro_Price', 'lead_time', 'Suggested_Price_15_Margin']
                            final_cols = [col for col in cols_to_show if col in df_sheet.columns]
                            df_sheet[final_cols].to_excel(writer, sheet_name=sheet_name, index=False)
                            format_excel_sheet(writer.sheets[sheet_name], df_sheet[final_cols])
                        else:
                            df_sheet.to_excel(writer, sheet_name=sheet_name, index=False)
                            format_excel_sheet(writer.sheets[sheet_name], df_sheet)
            
            self.log("----------------------------------------")
            self.log("SUCCESS: Analysis complete and report saved!")
            messagebox.showinfo("Success", f"Report has been saved successfully to:\n{save_path}")

        except Exception as e:
            self.log(f"ERROR: An error occurred: {e}")
            messagebox.showerror("Error", f"An unexpected error occurred during analysis:\n\n{e}")
        finally:
            self.run_button.config(state="normal")

if __name__ == "__main__":
    app = App()
    app.mainloop()