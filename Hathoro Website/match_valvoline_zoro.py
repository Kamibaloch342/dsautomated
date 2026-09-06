#!/usr/bin/env python3
"""
match_valvoline_zoro.py

Matches Amazon Keepa listings from Valvoline.xlsx against Zoro supplier catalog
in combined_reports.db using UPC, EAN, GTIN, and Manufacturer Part Number (MPN).

Outputs:
1. Valvoline_Zoro_Matched_Results.xlsx (Multi-tab formatted Excel report)
2. Valvoline_Zoro_Matched_Results.csv (Complete data CSV export)
"""

import os
import re
import sqlite3
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --- PATHS ---
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_INPUT_PATH = os.path.join(CURRENT_DIR, "Valvoline.xlsx")
DB_PATH = os.path.abspath(os.path.join(CURRENT_DIR, "..", "Automations", "TSV Files", "combined_reports.db"))
OUTPUT_EXCEL_PATH = os.path.join(CURRENT_DIR, "Valvoline_Zoro_Matched_Results.xlsx")
OUTPUT_CSV_PATH = os.path.join(CURRENT_DIR, "Valvoline_Zoro_Matched_Results.csv")


def extract_codes(val):
    """Extract clean digit strings from UPC / EAN / GTIN fields (splits multi-codes)."""
    if pd.isna(val) or val is None:
        return []
    parts = re.split(r'[,;\s/]+', str(val).strip())
    codes = []
    for p in parts:
        digits = re.sub(r'\D', '', p)
        if digits:
            codes.append(digits)
    return codes


def get_mpn_variants(val):
    """Generate normalized part number variants for resilient matching."""
    if pd.isna(val) or val is None:
        return []
    s = str(val).strip().upper()
    variants = {s}

    # Strip packaging and unit suffixes common in auto/lube listings
    cleaned = re.sub(r'[-_ ](EA|CS|FBA|\d+PK|CASE|PACK).*$', '', s)
    variants.add(cleaned)

    # Remove hyphens and whitespace
    no_hyphens = re.sub(r'[-_\s]', '', s)
    variants.add(no_hyphens)
    cleaned_no_hyphens = re.sub(r'[-_\s]', '', cleaned)
    variants.add(cleaned_no_hyphens)

    # Remove leading zeros if entirely numeric
    if cleaned_no_hyphens.isdigit():
        variants.add(cleaned_no_hyphens.lstrip('0'))

    return [v for v in variants if v]


def load_zoro_catalog(db_path):
    """Loads and indexes Zoro catalog items into memory for instant lookup."""
    print(f"Connecting to Zoro database: {db_path}...")
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()

    query = """
        SELECT zoro_no, manufacturer_part_number, upc_code, zoro_brand_name,
               short_description, long_description, sales_price, freight_surcharge,
               unit_price, level_one_name, level_two_name, is_dropship, lead_time,
               min_order_qty, max_order_qty, uom, availabilityStatus
        FROM zoro_data;
    """
    cur.execute(query)
    rows = cur.fetchall()
    conn.close()
    print(f"Loaded {len(rows):,} total Zoro catalog items.")

    upc_map = {}
    mpn_map = {}

    for r in rows:
        (z_no, z_mpn, z_upc, z_brand, z_desc, z_long_desc, z_price, z_freight,
         z_unit_price, z_l1, z_l2, z_drop, z_lead, z_min_qty, z_max_qty, z_uom, z_avail) = r

        item = {
            'zoro_no': z_no,
            'mpn': str(z_mpn).strip() if z_mpn else '',
            'upc': str(z_upc).strip() if z_upc else '',
            'brand': str(z_brand).strip() if z_brand else '',
            'short_desc': str(z_desc).strip() if z_desc else '',
            'long_desc': str(z_long_desc).strip() if z_long_desc else '',
            'sales_price': float(z_price) if z_price is not None else None,
            'freight_surcharge': float(z_freight) if z_freight is not None else 0.0,
            'unit_price': float(z_unit_price) if z_unit_price is not None else None,
            'level_one': z_l1 or '',
            'level_two': z_l2 or '',
            'is_dropship': z_drop,
            'lead_time': z_lead,
            'min_order_qty': z_min_qty,
            'max_order_qty': z_max_qty,
            'uom': z_uom or 'Each',
            'avail': z_avail or 'Unknown',
            'zoro_url': f"https://www.zoro.com/i/{z_no}/"
        }

        # Index by UPC / EAN / GTIN
        if z_upc:
            for c in extract_codes(z_upc):
                upc_map.setdefault(c, []).append(item)
                upc_map.setdefault(c.lstrip('0'), []).append(item)
                if len(c) == 12:
                    upc_map.setdefault('0' + c, []).append(item)
                    upc_map.setdefault('00' + c, []).append(item)

        # Index by MPN
        if z_mpn:
            for v in get_mpn_variants(z_mpn):
                mpn_map.setdefault(v, []).append(item)

    print("Zoro indexing complete.")
    return upc_map, mpn_map


def match_listings():
    upc_map, mpn_map = load_zoro_catalog(DB_PATH)

    print(f"\nReading Amazon listings from: {EXCEL_INPUT_PATH}...")
    df_raw = pd.read_excel(EXCEL_INPUT_PATH)
    df = df_raw.dropna(subset=['ASIN']).copy()
    print(f"Found {len(df)} Amazon listings to match.")

    all_records = []

    for idx, row in df.iterrows():
        asin = str(row['ASIN']).strip()
        amazon_title = str(row.get('Item Highlights') or '').strip()
        amazon_mfr_pn = str(row.get('Product Codes: PartNumber') or '').strip()
        upc_raw = row.get('Product Codes: UPC')
        ean_raw = row.get('Product Codes: EAN')
        gtin_raw = row.get('Product Codes: GTIN')
        amazon_url = str(row.get('URL: Amazon') or f"https://www.amazon.com/dp/{asin}")
        num_items = row.get('Number of Items')
        try:
            num_items = int(float(num_items)) if pd.notna(num_items) else 1
        except Exception:
            num_items = 1

        sales_rank = row.get('Sales Rank: Current')
        sales_rank_90d = row.get('Sales Rank: 90 days avg.')
        bought_past_month = row.get('Monthly Sales Trends: Bought in past month')
        reviews_rating = row.get('Reviews: Rating')
        reviews_count = row.get('Reviews: Rating Count')
        buy_box_seller = row.get('Buy Box: Buy Box Seller')
        fba_fee = row.get('FBA Pick&Pack Fee')
        referral_fee_pct = row.get('Referral Fee %')

        # Collect all candidate codes
        all_codes = []
        for val in [upc_raw, ean_raw, gtin_raw]:
            all_codes.extend(extract_codes(val))

        matched_zoro_items = []
        match_method = None

        # 1. First priority: Exact UPC / EAN / GTIN Match
        for c in all_codes:
            for cand in [c, c.lstrip('0')]:
                if cand in upc_map:
                    for z_item in upc_map[cand]:
                        # Avoid duplicates
                        if not any(m['zoro_no'] == z_item['zoro_no'] for m in matched_zoro_items):
                            matched_zoro_items.append(z_item)
                            match_method = 'UPC/EAN/GTIN Match'

        # 2. Second priority: MPN Match (with brand safety check)
        if not matched_zoro_items and amazon_mfr_pn:
            for v in get_mpn_variants(amazon_mfr_pn):
                if v in mpn_map:
                    # Verify brand is Valvoline or Zerex
                    cand_items = [
                        m for m in mpn_map[v]
                        if any(b in (m['brand'] or '').upper() for b in ['VALV', 'ZEREX'])
                    ]
                    for z_item in cand_items:
                        if not any(m['zoro_no'] == z_item['zoro_no'] for m in matched_zoro_items):
                            matched_zoro_items.append(z_item)
                            match_method = 'MFR Part# Match'

        # If multiple Zoro items matched via shared code, prefer the exact MPN match if available
        if len(matched_zoro_items) > 1 and amazon_mfr_pn:
            exact_mpn_cands = [
                m for m in matched_zoro_items
                if get_mpn_variants(m['mpn']) and any(v in get_mpn_variants(amazon_mfr_pn) for v in get_mpn_variants(m['mpn']))
            ]
            if exact_mpn_cands:
                matched_zoro_items = exact_mpn_cands

        if matched_zoro_items:
            for z_item in matched_zoro_items:
                zoro_price = z_item['sales_price']
                est_order_cost = (zoro_price * num_items) if zoro_price is not None else None
                pack_type = "Single Unit" if num_items == 1 else f"Multi-Pack ({num_items}x)"

                all_records.append({
                    'ASIN': asin,
                    'Match Status': 'Matched',
                    'Match Method': match_method,
                    'Zoro Item #': z_item['zoro_no'],
                    'Zoro MPN': z_item['mpn'],
                    'Zoro Brand': z_item['brand'],
                    'Zoro Description': z_item['short_desc'],
                    'Zoro Availability': z_item['avail'],
                    'Zoro Price ($)': zoro_price,
                    'Amazon Pack Qty': num_items,
                    'Zoro Est. Total Cost ($)': est_order_cost,
                    'Pack Type': pack_type,
                    'Zoro UOM': z_item['uom'],
                    'Amazon Part #': amazon_mfr_pn,
                    'Amazon UPC': upc_raw,
                    'Amazon Bought Past Month': bought_past_month,
                    'Amazon Sales Rank': sales_rank,
                    'Amazon Sales Rank (90d Avg)': sales_rank_90d,
                    'Reviews Rating': reviews_rating,
                    'Reviews Count': reviews_count,
                    'Buy Box Seller': buy_box_seller,
                    'FBA Pick&Pack Fee': fba_fee,
                    'Referral Fee %': referral_fee_pct,
                    'Zoro URL': z_item['zoro_url'],
                    'Amazon URL': amazon_url,
                    'Zoro Dropship Eligible': z_item['is_dropship'],
                    'Zoro Category L1': z_item['level_one'],
                    'Zoro Category L2': z_item['level_two']
                })
        else:
            all_records.append({
                'ASIN': asin,
                'Match Status': 'Not Found on Zoro',
                'Match Method': 'None',
                'Zoro Item #': '',
                'Zoro MPN': '',
                'Zoro Brand': '',
                'Zoro Description': 'Not available in Zoro catalog',
                'Zoro Availability': 'Not Found',
                'Zoro Price ($)': None,
                'Amazon Pack Qty': num_items,
                'Zoro Est. Total Cost ($)': None,
                'Pack Type': "Single Unit" if num_items == 1 else f"Multi-Pack ({num_items}x)",
                'Zoro UOM': '',
                'Amazon Part #': amazon_mfr_pn,
                'Amazon UPC': upc_raw,
                'Amazon Bought Past Month': bought_past_month,
                'Amazon Sales Rank': sales_rank,
                'Amazon Sales Rank (90d Avg)': sales_rank_90d,
                'Reviews Rating': reviews_rating,
                'Reviews Count': reviews_count,
                'Buy Box Seller': buy_box_seller,
                'FBA Pick&Pack Fee': fba_fee,
                'Referral Fee %': referral_fee_pct,
                'Zoro URL': '',
                'Amazon URL': amazon_url,
                'Zoro Dropship Eligible': '',
                'Zoro Category L1': '',
                'Zoro Category L2': ''
            })

    results_df = pd.DataFrame(all_records)
    print(f"\nMatching Summary:")
    print(f"Total Amazon Listings: {df['ASIN'].nunique()}")
    print(f"Matched ASINs: {results_df[results_df['Match Status'] == 'Matched']['ASIN'].nunique()}")
    print(f"Unmatched ASINs: {results_df[results_df['Match Status'] != 'Matched']['ASIN'].nunique()}")

    # Export CSV
    results_df.to_csv(OUTPUT_CSV_PATH, index=False)
    print(f"Saved raw CSV to: {OUTPUT_CSV_PATH}")

    # Build formatted Multi-Tab Excel
    export_styled_excel(results_df, OUTPUT_EXCEL_PATH)
    print(f"Saved styled Excel report to: {OUTPUT_EXCEL_PATH}")


def export_styled_excel(df, file_path):
    """Exports structured, beautifully formatted multi-tab Excel report."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    accent_blue_fill = PatternFill(start_color="2B547E", end_color="2B547E", fill_type="solid")
    kpi_card_fill = PatternFill(start_color="EAF2F8", end_color="EAF2F8", fill_type="solid")
    kpi_header_fill = PatternFill(start_color="2471A3", end_color="2471A3", fill_type="solid")
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    kpi_title_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    kpi_val_font = Font(name="Calibri", size=18, bold=True, color="1B365D")
    data_font = Font(name="Calibri", size=10)
    bold_data_font = Font(name="Calibri", size=10, bold=True)
    link_font = Font(name="Calibri", size=10, color="0000EE", underline="single")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    kpi_border = Border(
        left=Side(style='medium', color='2471A3'),
        right=Side(style='medium', color='2471A3'),
        top=Side(style='medium', color='2471A3'),
        bottom=Side(style='medium', color='2471A3')
    )

    in_stock_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    in_stock_font = Font(name="Calibri", size=10, color="155724", bold=True)

    out_stock_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    out_stock_font = Font(name="Calibri", size=10, color="721C24", bold=True)

    limited_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    limited_font = Font(name="Calibri", size=10, color="856404", bold=True)

    # --- TAB 1: SUMMARY DASHBOARD ---
    ws_summary = wb.create_sheet(title="Executive Summary")
    ws_summary.views.sheetView[0].showGridLines = True

    # Title
    ws_summary["A1"] = "HATHORO - VALVOLINE AMAZON TO ZORO MATCHING REPORT"
    ws_summary["A1"].font = Font(name="Calibri", size=16, bold=True, color="1B365D")
    ws_summary["A2"] = f"Generated on {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')} | Source Database: combined_reports.db (Zoro Partner Feed)"
    ws_summary["A2"].font = Font(name="Calibri", size=10, italic=True, color="555555")

    total_listings = df['ASIN'].nunique()
    matched_listings = df[df['Match Status'] == 'Matched']['ASIN'].nunique()
    unmatched_listings = total_listings - matched_listings
    match_pct = (matched_listings / total_listings * 100) if total_listings else 0

    matched_sub = df[df['Match Status'] == 'Matched']
    in_stock_count = matched_sub[matched_sub['Zoro Availability'] == 'In Stock']['ASIN'].nunique()
    out_stock_count = matched_sub[matched_sub['Zoro Availability'] == 'Out of Stock']['ASIN'].nunique()
    limited_count = matched_sub[matched_sub['Zoro Availability'] == 'Limited Stock']['ASIN'].nunique()
    multi_pack_count = matched_sub[matched_sub['Amazon Pack Qty'] > 1]['ASIN'].nunique()

    kpis = [
        ("Total Keepa Listings", f"{total_listings}", "B4", "C5"),
        ("Matched in Zoro", f"{matched_listings} ({match_pct:.1f}%)", "D4", "E5"),
        ("In Stock on Zoro", f"{in_stock_count}", "F4", "G5"),
        ("Out of Stock / Limited", f"{out_stock_count + limited_count}", "H4", "I5"),
    ]

    for title, val, top_left, bot_right in kpis:
        col1 = top_left[0]
        row1 = int(top_left[1])
        col2 = bot_right[0]
        row2 = int(bot_right[1])
        ws_summary.merge_cells(f"{top_left}:{col2}{row1}")
        ws_summary.merge_cells(f"{col1}{row2}:{bot_right}")
        
        c_title = ws_summary[f"{col1}{row1}"]
        c_title.value = title
        c_title.font = kpi_title_font
        c_title.fill = kpi_header_fill
        c_title.alignment = Alignment(horizontal="center", vertical="center")

        c_val = ws_summary[f"{col1}{row2}"]
        c_val.value = val
        c_val.font = kpi_val_font
        c_val.fill = kpi_card_fill
        c_val.alignment = Alignment(horizontal="center", vertical="center")

    # Table of Breakdown
    ws_summary["B8"] = "Metric"
    ws_summary["B8"].font = header_font
    ws_summary["B8"].fill = navy_header_fill
    ws_summary["C8"] = "Count / Value"
    ws_summary["C8"].font = header_font
    ws_summary["C8"].fill = navy_header_fill
    ws_summary["D8"] = "Notes / Dropshipping Implication"
    ws_summary["D8"].font = header_font
    ws_summary["D8"].fill = navy_header_fill

    breakdown_rows = [
        ("Total Products Analyzed", total_listings, "All unique ASINs from Valvoline.xlsx Keepa export"),
        ("Matched Products in Zoro", matched_listings, f"{match_pct:.1f}% match rate via exact UPC/EAN and MPN"),
        ("Ready to Sell (In Stock)", in_stock_count, "Products actively available in Zoro inventory for fulfillment"),
        ("Multi-Pack / Bundle Listings", multi_pack_count, "Listings sold on Amazon as 4-pack, 6-pack, or 10-pack (bundle multiplier needed)"),
        ("Out of Stock on Zoro", out_stock_count, "Keep monitored for restocking"),
        ("Limited Stock on Zoro", limited_count, "Stock quantity is constrained"),
        ("Unmatched / Not Distributed", unmatched_listings, "Retail-only 5-Qt jugs or discontinued OEM SKUs not in Zoro partner feed"),
    ]

    for idx, (metric, val, note) in enumerate(breakdown_rows, start=9):
        ws_summary[f"B{idx}"] = metric
        ws_summary[f"B{idx}"].font = bold_data_font
        ws_summary[f"B{idx}"].border = thin_border

        ws_summary[f"C{idx}"] = val
        ws_summary[f"C{idx}"].font = data_font
        ws_summary[f"C{idx}"].border = thin_border
        ws_summary[f"C{idx}"].alignment = Alignment(horizontal="center")

        ws_summary[f"D{idx}"] = note
        ws_summary[f"D{idx}"].font = data_font
        ws_summary[f"D{idx}"].border = thin_border

    # Adjust summary column widths
    ws_summary.column_dimensions['A'].width = 5
    ws_summary.column_dimensions['B'].width = 30
    ws_summary.column_dimensions['C'].width = 18
    ws_summary.column_dimensions['D'].width = 65
    ws_summary.column_dimensions['E'].width = 12
    ws_summary.column_dimensions['F'].width = 18
    ws_summary.column_dimensions['G'].width = 12
    ws_summary.column_dimensions['H'].width = 18
    ws_summary.column_dimensions['I'].width = 12

    # --- HELPER FUNCTION TO POPULATE DATA TABS ---
    def populate_sheet(ws, sub_df, title):
        ws.views.sheetView[0].showGridLines = True
        columns = [
            ("ASIN", 14),
            ("Match Status", 14),
            ("Match Method", 18),
            ("Zoro Item #", 14),
            ("Zoro MPN", 16),
            ("Zoro Brand", 14),
            ("Zoro Description", 38),
            ("Zoro Availability", 16),
            ("Zoro Price ($)", 14),
            ("Amazon Pack Qty", 15),
            ("Zoro Est. Total Cost ($)", 22),
            ("Pack Type", 15),
            ("Zoro UOM", 12),
            ("Amazon Part #", 16),
            ("Amazon UPC", 16),
            ("Amazon Bought Past Month", 24),
            ("Amazon Sales Rank", 18),
            ("Zoro URL", 32),
            ("Amazon URL", 32)
        ]

        # Header Row
        for col_idx, (col_name, col_width) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = navy_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col_idx)].width = col_width

        ws.row_dimensions[1].height = 28

        # Data Rows
        for row_idx, r in enumerate(sub_df.to_dict(orient='records'), start=2):
            ws.row_dimensions[row_idx].height = 20
            for col_idx, (col_name, _) in enumerate(columns, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                val = r.get(col_name)
                cell.border = thin_border
                cell.font = data_font

                if col_name in ["Zoro Price ($)", "Zoro Est. Total Cost ($)"]:
                    cell.value = val
                    if val is not None:
                        cell.number_format = '$#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif col_name in ["Amazon Bought Past Month", "Amazon Sales Rank", "Amazon Pack Qty"]:
                    cell.value = val
                    if val is not None and not pd.isna(val):
                        cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="right")
                elif col_name == "Zoro Availability":
                    cell.value = val
                    cell.alignment = Alignment(horizontal="center")
                    if val == "In Stock":
                        cell.fill = in_stock_fill
                        cell.font = in_stock_font
                    elif val == "Out of Stock":
                        cell.fill = out_stock_fill
                        cell.font = out_stock_font
                    elif val == "Limited Stock":
                        cell.fill = limited_fill
                        cell.font = limited_font
                elif col_name == "Match Status":
                    cell.value = val
                    cell.alignment = Alignment(horizontal="center")
                    if val == "Matched":
                        cell.fill = in_stock_fill
                        cell.font = in_stock_font
                    else:
                        cell.fill = out_stock_fill
                        cell.font = out_stock_font
                elif col_name in ["Zoro URL", "Amazon URL"]:
                    if val and str(val).startswith("http"):
                        cell.value = f'=HYPERLINK("{val}", "View on {col_name.split()[0]}")'
                        cell.font = link_font
                        cell.alignment = Alignment(horizontal="center")
                    else:
                        cell.value = "-"
                        cell.alignment = Alignment(horizontal="center")
                else:
                    cell.value = val if pd.notna(val) else ""
                    if col_name in ["ASIN", "Zoro Item #", "Zoro MPN", "Amazon Part #", "Pack Type", "Zoro UOM"]:
                        cell.alignment = Alignment(horizontal="center")

    # --- TAB 2: IN-STOCK MATCHED (READY TO SELL) ---
    ws_instock = wb.create_sheet(title="In-Stock (Ready to List)")
    df_instock = df[(df['Match Status'] == 'Matched') & (df['Zoro Availability'] == 'In Stock')].copy()
    df_instock.sort_values(by="Amazon Bought Past Month", ascending=False, inplace=True)
    populate_sheet(ws_instock, df_instock, "In-Stock Zoro Products")

    # --- TAB 3: ALL MATCHED PRODUCTS ---
    ws_matched = wb.create_sheet(title="All Matched Products")
    df_matched = df[df['Match Status'] == 'Matched'].copy()
    df_matched.sort_values(by="Amazon Bought Past Month", ascending=False, inplace=True)
    populate_sheet(ws_matched, df_matched, "All Matched Products")

    # --- TAB 4: MULTI-PACKS / BUNDLES ---
    ws_packs = wb.create_sheet(title="Multi-Pack Bundles")
    df_packs = df[(df['Match Status'] == 'Matched') & (df['Amazon Pack Qty'] > 1)].copy()
    populate_sheet(ws_packs, df_packs, "Multi-Pack Listings")

    # --- TAB 5: UNMATCHED LISTINGS ---
    ws_unmatched = wb.create_sheet(title="Unmatched (Not on Zoro)")
    df_unmatched = df[df['Match Status'] != 'Matched'].copy()
    df_unmatched.sort_values(by="Amazon Bought Past Month", ascending=False, inplace=True)
    populate_sheet(ws_unmatched, df_unmatched, "Unmatched Products")

    wb.save(file_path)


if __name__ == "__main__":
    match_listings()
