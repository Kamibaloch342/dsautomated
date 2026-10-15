import os
import re
import json
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = "/Users/kamibaloch342/Downloads/Dropshipping/Development_Unit/Hathoro Website"
DRUMS_DIR = os.path.join(BASE_DIR, "Drums")
CURATED_PATH = os.path.join(DRUMS_DIR, "curated_titles.json")
SUMMARY_PATH = os.path.join(DRUMS_DIR, "drums_processing_summary.json")
CSV_PATH = os.path.join(DRUMS_DIR, "Drums_Zoro_Matched_Results.csv")
EXCEL_PATH = os.path.join(DRUMS_DIR, "Drums_Zoro_Matched_Results.xlsx")

def main():
    print("Loading curated human titles...")
    with open(CURATED_PATH, "r", encoding="utf-8") as f:
        curated = json.load(f)
    print(f"Loaded {len(curated)} curated titles.")

    # 1. Update listing copy, details, and data.json for each product
    updated_copy = 0
    updated_details = 0
    updated_data = 0

    for sku, info in curated.items():
        p_title = info["primary_title"]
        s_title = info["short_title"]
        sku_dir = os.path.join(DRUMS_DIR, sku)

        # A. Listing Copy
        copy_file = os.path.join(sku_dir, f"{sku}_listing_copy.txt")
        if os.path.exists(copy_file):
            with open(copy_file, "r", encoding="utf-8") as f:
                content = f.read()

            content = re.sub(r"Product Name:\s*.*", f"Product Name: {p_title}", content, count=1)
            s1_pattern = r"(\[RECOMMENDED PRIMARY TITLE.*?\]\n)[^\n]+(\n\n\[CLEAN SHORT TITLE.*?\]\n)[^\n]+"
            def s1_repl(m):
                return f"{m.group(1)}{p_title}{m.group(2)}{s_title}"
            content = re.sub(s1_pattern, s1_repl, content)

            with open(copy_file, "w", encoding="utf-8") as f:
                f.write(content)
            updated_copy += 1

        # B. Details Text
        details_file = os.path.join(sku_dir, f"{sku}_details.txt")
        if os.path.exists(details_file):
            with open(details_file, "r", encoding="utf-8") as f:
                content = f.read()

            content = re.sub(r"Title:\s*.*", f"Title: {p_title}", content, count=1)
            # Add Short Title if not present
            if "Short Title:" not in content:
                content = content.replace(f"Title: {p_title}\n", f"Title: {p_title}\nShort Title: {s_title}\n")
            else:
                content = re.sub(r"Short Title:\s*.*", f"Short Title: {s_title}", content, count=1)

            with open(details_file, "w", encoding="utf-8") as f:
                f.write(content)
            updated_details += 1

        # C. Data JSON
        data_file = os.path.join(sku_dir, f"{sku}_data.json")
        if os.path.exists(data_file):
            with open(data_file, "r", encoding="utf-8") as f:
                data = json.load(f)

            if "original_title" not in data:
                data["original_title"] = data.get("title", "")
            data["title"] = p_title
            data["short_title"] = s_title

            with open(data_file, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2)
            updated_data += 1

    print(f"Updated {updated_copy} listing_copy files.")
    print(f"Updated {updated_details} details files.")
    print(f"Updated {updated_data} data.json files.")

    # 2. Update drums_processing_summary.json
    with open(SUMMARY_PATH, "r", encoding="utf-8") as f:
        summary_items = json.load(f)

    for item in summary_items:
        sku = item.get("zoroNo")
        if sku in curated:
            if "original_title" not in item:
                item["original_title"] = item.get("title", "")
            item["title"] = curated[sku]["primary_title"]
            item["short_title"] = curated[sku]["short_title"]

    with open(SUMMARY_PATH, "w", encoding="utf-8") as f:
        json.dump(summary_items, f, indent=2)
    print(f"Updated {len(summary_items)} items in summary JSON.")

    # 3. Update CSV export
    csv_headers = [
        "zoroNo", "title", "short_title", "brand", "mfrNo", "price", "rating", "reviews",
        "category_type", "images_saved", "status", "url", "capacity", "material"
    ]
    with open(CSV_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=csv_headers, extrasaction="ignore")
        writer.writeheader()
        for item in summary_items:
            writer.writerow(item)
    print(f"Updated CSV export with {len(summary_items)} rows.")

    # 4. Rebuild Excel Workbook
    drum_items = [it for it in summary_items if it.get("category_type") != "Industrial Cleaner / Degreaser"]
    cleaning_items = [it for it in summary_items if it.get("category_type") == "Industrial Cleaner / Degreaser"]

    wb = openpyxl.Workbook()
    FONT_NAME = "Segoe UI"
    NAVY = "1B365D"
    WHITE = "FFFFFF"
    ZEBRA = "F4F7FB"
    BORDER_COLOR = "D9D9D9"

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    def write_sheet(ws, banner_title, rows_data):
        ws.views.sheetView[0].showGridLines = True
        ws.merge_cells("A1:L1")
        tcell = ws["A1"]
        tcell.value = banner_title
        tcell.font = Font(name=FONT_NAME, size=14, bold=True, color=WHITE)
        tcell.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
        tcell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        headers = [
            "Item #", "Zoro SKU", "Brand", "Curated Primary Title", "Clean Short Title",
            "Category / Type", "Capacity / Size", "Material / Formula", "Price ($)", "Rating (★)", "Reviews", "Zoro Link"
        ]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(3, col_idx, h)
            cell.font = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
            cell.fill = PatternFill(start_color="335588", end_color="335588", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[3].height = 24

        for row_idx, r in enumerate(rows_data, 4):
            is_even = (row_idx % 2 == 0)
            row_fill = PatternFill(start_color=ZEBRA, end_color=ZEBRA, fill_type="solid") if is_even else PatternFill(fill_type=None)

            ws.cell(row_idx, 1, row_idx - 3).alignment = Alignment(horizontal="center")
            ws.cell(row_idx, 2, r.get('zoroNo', '')).alignment = Alignment(horizontal="center")
            ws.cell(row_idx, 3, r.get('brand', 'Zoro Select')).alignment = Alignment(horizontal="left")
            ws.cell(row_idx, 4, r.get('title', '')).alignment = Alignment(horizontal="left")
            ws.cell(row_idx, 5, r.get('short_title', '')).alignment = Alignment(horizontal="left")
            ws.cell(row_idx, 6, r.get('category_type', '')).alignment = Alignment(horizontal="center")
            ws.cell(row_idx, 7, r.get('capacity', 'N/A')).alignment = Alignment(horizontal="center")
            ws.cell(row_idx, 8, r.get('material', 'N/A')).alignment = Alignment(horizontal="center")

            p_cell = ws.cell(row_idx, 9, r.get('price', 0.0))
            p_cell.number_format = "$#,##0.00"
            p_cell.alignment = Alignment(horizontal="right")

            r_cell = ws.cell(row_idx, 10, r.get('rating', 5.0))
            r_cell.number_format = "0.0"
            r_cell.alignment = Alignment(horizontal="center")

            c_cell = ws.cell(row_idx, 11, r.get('reviews', 1))
            c_cell.number_format = "#,##0"
            c_cell.alignment = Alignment(horizontal="center")

            url_cell = ws.cell(row_idx, 12, f"https://www.zoro.com/i/{r.get('zoroNo')}/")
            url_cell.font = Font(name=FONT_NAME, size=9, color="0000FF", underline="single")
            url_cell.alignment = Alignment(horizontal="left")

            for c in range(1, 13):
                cell = ws.cell(row_idx, c)
                if c != 12:
                    cell.font = Font(name=FONT_NAME, size=9)
                if row_fill.fill_type:
                    cell.fill = row_fill
                cell.border = thin_border
            ws.row_dimensions[row_idx].height = 20

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 60)

    # Sheet 1: Drums (100)
    ws1 = wb.active
    ws1.title = "Ready to List - Drums (100)"
    write_sheet(ws1, "HATHORO E-COMMERCE: 100 VERIFIED DRUM CATALOG LISTINGS", drum_items)

    # Sheet 2: Cleaning (8)
    ws2 = wb.create_sheet(title="Cleaning Products (8)")
    write_sheet(ws2, "HATHORO E-COMMERCE: VERIFIED INDUSTRIAL CLEANING CATALOG LISTINGS", cleaning_items)

    wb.save(EXCEL_PATH)
    print(f"Successfully generated updated Excel workbook at: {EXCEL_PATH}")

if __name__ == "__main__":
    main()
