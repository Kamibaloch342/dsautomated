# -*- coding: utf-8 -*-
"""
Professional Web Table Scraper with Hybrid Pagination and Advanced Formatting.

VERSION: 8.0 (Final)

This script uses a persistent Chrome profile to scrape a data table. It offers
both automatic and manual pagination control and saves a professionally formatted
Excel file with a dynamic name based on the target URL.

WORKFLOW (v8.0):
1.  Launches Chrome, connects Selenium, and navigates to the URL.
2.  Performs a stabilization maneuver to ensure the first page loads correctly.
3.  Scrapes the first page and then prompts the user for the next action.
4.  User can choose (A)utomatic, (M)anual, or (S)top.
5.  The script continues based on user input until stopped.
6.  It combines all data, cleans it, and saves it to a dynamically named Excel file.
7.  Finally, it applies advanced formatting to the Excel file (borders, no gridlines, etc.).
"""

import os
import sys
import time
import subprocess
import pandas as pd
from collections import Counter
from urllib.parse import urlparse, parse_qs
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.options import Options as ChromeOptions
from selenium.common.exceptions import TimeoutException, WebDriverException, NoSuchElementException

# --- 🛠️ CONFIGURATION ---
TARGET_URL = "https://app.smartscout.com/app/sellers?id=1494364&tab=1&label=Beyond%20Box"
DEBUGGING_PORT = 9224

# --- 📂 PROFILE SETUP ---
user_profile_dir = os.path.expanduser("~")
if sys.platform == "darwin":  # macOS
    profile_path = os.path.join(user_profile_dir, "Library", "Application Support", "Google", "Chrome", "smartscout_extractor")
    CHROME_EXECUTABLE_PATH = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
else:  # Windows
    profile_path = os.path.join(user_profile_dir, "AppData", "Local", "Google", "Chrome", "User Data", "smartscout_extractor")
    CHROME_EXECUTABLE_PATH = "C:\\Program Files\\Google\\Chrome\\Application\\chrome.exe"
    if not os.path.exists(CHROME_EXECUTABLE_PATH):
        CHROME_EXECUTABLE_PATH = "C:\\Program Files (x86)\\Google\\Chrome\\Application\\chrome.exe"

print(f"🖥️  Using professional profile path: {profile_path}")

# --- EXCEL FORMATTING FUNCTION ---
def format_excel_file(filename, df):
    """Applies advanced formatting to the saved Excel file."""
    try:
        import openpyxl
        from openpyxl.styles import Font, Border, Side
        from openpyxl.utils import get_column_letter

        print(f"\n🎨 Applying advanced formatting to '{filename}'...")

        workbook = openpyxl.load_workbook(filename)
        worksheet = workbook.active

        # 1. Define styles
        font_style = Font(name='Calibri', size=11)
        thin_border = Border(left=Side(style='thin'),
                             right=Side(style='thin'),
                             top=Side(style='thin'),
                             bottom=Side(style='thin'))

        # 2. Apply styles and borders to all cells
        for row in worksheet.iter_rows():
            for cell in row:
                cell.font = font_style
                cell.border = thin_border

        # 3. Format UPC column as a number with zero decimals
        if 'UPC' in df.columns:
            upc_col_index = df.columns.get_loc('UPC') + 1
            upc_col_letter = get_column_letter(upc_col_index)
            for cell in worksheet[upc_col_letter]:
                if cell.row > 1: # Skip header
                    cell.number_format = '0'

        # 4. Auto-fit column widths
        for col in worksheet.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 2)
            worksheet.column_dimensions[column].width = adjusted_width

        # 5. Remove gridlines
        worksheet.sheet_view.showGridLines = False

        workbook.save(filename)
        print("✅ Excel formatting complete.")

    except ImportError:
        print("\n⚠️ 'openpyxl' is not installed. Skipping advanced Excel formatting.")
        print("   To enable formatting, run: pip install openpyxl")
    except Exception as e:
        print(f"\n❌ An error occurred during Excel formatting: {e}")


# --- ▶️ MAIN EXECUTION ---
driver = None
chrome_process = None
try:
    # --- Dynamic Filename ---
    parsed_url = urlparse(TARGET_URL)
    query_params = parse_qs(parsed_url.query)
    seller_name = query_params.get('label', ['scraped_data'])[0].replace('+', '')
    OUTPUT_FILE_NAME = f"{seller_name}.xlsx"

    # 1. Launch a quiet Chrome
    print("🚀 Launching persistent, quiet Chrome with remote debugging...")
    command = (f'"{CHROME_EXECUTABLE_PATH}" '
               f'--remote-debugging-port={DEBUGGING_PORT} '
               f'--user-data-dir="{profile_path}" '
               '--log-level=3 '
               '--disable-logging')
    chrome_process = subprocess.Popen(command, shell=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print(f"✅ Chrome process started with PID: {chrome_process.pid}. Waiting a moment to initialize...")
    time.sleep(5)

    # 2. Connect Selenium
    print("🔗 Connecting Selenium to the running Chrome instance...")
    chrome_options = ChromeOptions()
    chrome_options.add_experimental_option("debuggerAddress", f"127.0.0.1:{DEBUGGING_PORT}")
    driver = webdriver.Chrome(options=chrome_options)
    print("✅ Selenium connected successfully.")

    # 3. Navigate directly to the target URL
    print("-" * 50)
    print(f"🚗 Navigating to: {TARGET_URL}")
    driver.get(TARGET_URL)
    print("\nPage navigation initiated. Starting the scraping process...")
    
    all_page_dataframes = []
    page_number = 1
    scrape_mode = "initial" # Modes: initial, automatic, manual

    # --- Stabilize Page 1 ---
    try:
        print("  -> Initializing page state... Waiting for controls.")
        wait = WebDriverWait(driver, 20)
        next_button_stabilize = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, "div[aria-label='Next Page']")))
        
        if next_button_stabilize.get_attribute('aria-disabled') == 'false':
            print("  -> Stabilizing page 1: Navigating to page 2...")
            driver.execute_script("arguments[0].click();", next_button_stabilize)
            time.sleep(5)
            prev_button_stabilize = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, "div[aria-label='Previous Page']")))
            print("  -> Stabilizing page 1: Navigating back to page 1...")
            driver.execute_script("arguments[0].click();", prev_button_stabilize)
            time.sleep(5)
            print("  -> Page 1 stabilized and is ready for scraping.")
        else:
            print("  -> Only one page detected. Skipping stabilization.")
    except Exception:
        print("  -> Single page or pagination controls not found immediately. Proceeding directly.")
    
    print("-" * 50)

    # 4. Main scraping loop
    while True:
        try:
            print(f"📖 Scraping page {page_number}...")
            wait = WebDriverWait(driver, 40)
            wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, "div[role='grid']")))
            print("  -> Data grid container is present.")
            print("  -> Waiting 10 seconds for all data rows to fully load...")
            time.sleep(10)
            
            grid_container = driver.find_element(By.CSS_SELECTOR, "div[role='grid']")
            headers = [h.text.strip() for h in grid_container.find_elements(By.CSS_SELECTOR, "div[role='columnheader']") if h.text.strip()]
            rows = grid_container.find_elements(By.CSS_SELECTOR, "div[role='row']")
            all_rows_data = [[cell.text.strip() for cell in r.find_elements(By.CSS_SELECTOR, "div[role='gridcell']")] for r in rows[1:]]
            
            if all_rows_data:
                row_lengths = Counter(len(r) for r in all_rows_data)
                common_len = row_lengths.most_common(1)[0][0]
                if common_len > len(headers):
                    missing = common_len - len(headers)
                    print(f"  -> Aligning data by prepending {missing} header(s).")
                    for i in range(missing): headers.insert(i, f'UNKNOWN_COL_{i+1}')
                    headers[0] = 'UPC'
                
                valid_rows = [r for r in all_rows_data if len(r) == len(headers)]
                if valid_rows:
                    df = pd.DataFrame(valid_rows, columns=headers)
                    all_page_dataframes.append(df)
                    print(f"  -> Successfully scraped {len(df)} rows.")
                else: print("  -> ⚠️ No valid rows found after alignment.")
            else: print("  -> ⚠️ No data rows found on this page.")

            # --- HYBRID PAGINATION LOGIC ---
            if scrape_mode == 'automatic':
                # Continue with automatic clicking
                pass
            else: # initial or manual mode
                print("\n" + "="*25)
                user_input = input("Choose next action: (A)utomatic, (M)anual, or (S)top & Save: ").strip().lower()
                print("="*25 + "\n")

                if user_input == 's':
                    print("✅ Stopping script as requested.")
                    break
                elif user_input == 'a':
                    scrape_mode = 'automatic'
                    print("...Switching to Automatic mode...")
                elif user_input == 'm':
                    scrape_mode = 'manual'
                    input("...Please navigate to the desired page, then press Enter to continue...")
                    page_number += 1
                    time.sleep(5)
                    continue # Skip automatic click and restart loop for the new page
                else:
                    print("⚠️ Invalid input. Stopping script.")
                    break

            # Automatic click logic
            try:
                next_button = driver.find_element(By.CSS_SELECTOR, "div[aria-label='Next Page']")
                if next_button.get_attribute('aria-disabled') == 'true':
                    print("\n✅ Reached the last page.")
                    break
                print("  -> Automatically clicking 'Next Page' button...")
                driver.execute_script("arguments[0].click();", next_button)
                page_number += 1
                time.sleep(5) 
            except NoSuchElementException:
                print("\n✅ No 'Next Page' button found. Stopping.")
                break

        except TimeoutException:
            print("\n⚠️ Timed out waiting for the page to load. Stopping scrape.")
            break
        except Exception as e:
            print(f"\n❌ An unexpected error occurred during scraping: {e}")
            break

    # 5. Combine, clean, and save
    if all_page_dataframes:
        print("\n" + "=" * 50)
        print(" consolidating all scraped data...")
        final_df = pd.concat(all_page_dataframes, ignore_index=True)
        print(f"  -> Total rows scraped: {len(final_df)}")

        print("\n🧹 Cleaning the final dataset...")
        cols_to_clean = {'Estimated Monthly Revenue': ['$', ','], 'Offer Price': ['$', ','], 'Buy Box Percentage': ['%'], 'Rank': [',']}
        for col, chars in cols_to_clean.items():
            if col in final_df.columns:
                final_df[col] = final_df[col].astype(str).str.replace(r'[^\d.]', '', regex=True)
                final_df[col] = pd.to_numeric(final_df[col], errors='coerce').fillna(0)
        print("✅ Data cleaning complete.")

        # Save the initial Excel file
        script_dir = os.path.dirname(os.path.abspath(__file__))
        output_file_path = os.path.join(script_dir, OUTPUT_FILE_NAME)
        print(f"\n💾 Saving data to '{output_file_path}'...")
        final_df.to_excel(output_file_path, index=False)
        print("🎉 Success! Initial file has been saved.")

        # Apply advanced formatting to the saved file
        format_excel_file(output_file_path, final_df)
        print("=" * 50)
    else:
        print("\nNo data was scraped. The script will now exit.")

except WebDriverException as e:
    print(f"\n❌ A connection error occurred: {e}")
except Exception as e:
    print(f"\n❌ A fatal error occurred: {e}")

finally:
    # 6. Cleanup
    if chrome_process and chrome_process.poll() is None:
        print("🔪 Terminating the Chrome process.")
        chrome_process.terminate()
    print("🎉 Script finished.")